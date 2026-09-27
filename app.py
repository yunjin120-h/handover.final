import hashlib
import io
import inspect
import json
import re
import time
import traceback
import uuid
from datetime import date, datetime
from pathlib import Path

import pandas as pd
import streamlit as st
import streamlit.components.v1 as components
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt

import ai_engine as ai
from memo_parser import parse_memo_text


# =========================================================
# 팀원 일정표 모듈 - HTML/JavaScript 원본 내장
#  (AI 버전: Python이 AI_* 상수에 문서 분석 결과를 주입합니다)
# =========================================================
ONBOARDING_SCHEDULE_HTML = r'''<!DOCTYPE html>
<html lang="ko">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>후임자 온보딩 일정 설계</title>
<link rel="preconnect" href="https://fonts.googleapis.com">
<link href="https://fonts.googleapis.com/css2?family=Space+Grotesk:wght@500;600;700&family=IBM+Plex+Sans+KR:wght@400;500;600;700&family=IBM+Plex+Mono:wght@400;500&display=swap" rel="stylesheet">
<script src="https://cdnjs.cloudflare.com/ajax/libs/mammoth/1.6.0/mammoth.browser.min.js"></script>
<style>
  :root{
    --ink:#161A1F; --paper:#F4F7FB; --surface:#FFFFFF;
    --line:#DCE5F0; --line-strong:#BFCDE0;
    --accent:#2F6FED; --accent-dark:#1E4FAE; --accent-soft:#EAF1FF;
    --amber:#9A6B14; --amber-soft:#F6E9D2;
    --warn:#B14A16; --warn-soft:#F5E2D5;
    --muted:#64748B;
    --shadow: 0 1px 2px rgba(20,20,15,0.04), 0 8px 24px rgba(20,20,15,0.05);
  }
  *{box-sizing:border-box;}
  html,body{margin:0;padding:0;}
  body{background:var(--paper); color:var(--ink); font-family:'IBM Plex Sans KR', sans-serif; line-height:1.55; -webkit-font-smoothing:antialiased;}
  .mono{font-family:'IBM Plex Mono', monospace;}
  header{border-bottom:1px solid var(--line); background:var(--surface); padding:22px 32px;}
  header .mark{width:34px;height:34px;border-radius:8px;background:var(--accent);display:inline-flex;align-items:center;justify-content:center;color:white;font-weight:700;font-size:14px;font-family:'Space Grotesk',sans-serif;margin-right:10px;vertical-align:middle;}
  header h1{font-size:18px;margin:0;font-weight:700;display:inline;vertical-align:middle;}
  header p{margin:6px 0 0;font-size:12.5px;color:var(--muted);}
  .layout{display:grid; grid-template-columns:320px 1fr; gap:20px; padding:24px; max-width:1360px; margin:0 auto; align-items:start;}
  @media (max-width:980px){ .layout{grid-template-columns:1fr;} }
  .panel{background:var(--surface); border:1px solid var(--line); border-radius:14px; box-shadow:var(--shadow);}
  .panel-head{padding:16px 18px 12px; border-bottom:1px solid var(--line);}
  .panel-head h2{font-size:13px;margin:0;font-weight:700;text-transform:uppercase;letter-spacing:0.05em;color:var(--muted);}
  .panel-body{padding:14px 18px 18px;}
  .source-block{margin-bottom:14px;}
  .source-block:last-child{margin-bottom:0;}
  .source-label{font-size:12px;font-weight:600;margin-bottom:5px;}
  .source-label .req{font-size:10.5px;color:var(--muted);font-weight:400;}
  .dropzone{border:1.5px dashed var(--line-strong); border-radius:10px; padding:12px 8px; text-align:center; cursor:pointer; background:#FBFBF9;}
  .dropzone.drag{border-color:var(--accent); background:var(--accent-soft);}
  .dropzone .dz-title{font-size:11.8px;font-weight:600;}
  .dropzone .dz-sub{font-size:10.4px;color:var(--muted);margin-top:2px;}
  input[type=file]{display:none;}
  .file-list{margin-top:6px;display:flex;flex-direction:column;gap:4px;}
  .file-item{display:flex;align-items:center;gap:6px;font-size:11px;padding:5px 8px;border-radius:6px;background:#FBFBF9;border:1px solid var(--line);}
  .file-item .fname{flex:1;overflow:hidden;text-overflow:ellipsis;white-space:nowrap;font-weight:500;}
  .file-item .fstatus{font-size:9.5px;color:var(--muted);font-family:'IBM Plex Mono',monospace;}
  .file-item .fstatus.ok{color:var(--accent-dark);}
  .file-item .fstatus.err{color:var(--warn);}
  .file-item .fremove{cursor:pointer;color:var(--muted);font-size:12px;padding:1px 3px;}
  .btn{appearance:none;border:none;cursor:pointer;font-family:'IBM Plex Sans KR',sans-serif;font-weight:600;font-size:13px;border-radius:8px;padding:9px 14px;}
  .btn:disabled{opacity:0.5;cursor:default;}
  .btn-primary{background:var(--accent); color:white;}
  .btn-primary:hover:not(:disabled){background:var(--accent-dark);}
  .btn-ghost{background:transparent;color:var(--accent-dark);border:1px solid var(--line-strong);}
  .btn-block{width:100%;text-align:center;}
  .row-btns{display:flex;gap:8px;margin-top:10px;flex-wrap:wrap;}
  .week-tabs{display:flex;gap:8px;margin-bottom:16px;flex-wrap:wrap;}
  .week-tab{border:1px solid var(--line-strong);border-radius:10px;padding:8px 14px;cursor:pointer;background:var(--surface);font-size:12.5px;font-weight:600;}
  .week-tab.active{border-color:var(--accent);background:var(--accent-soft);}
  .grid-wrap{overflow-x:auto;}
  table.sched{border-collapse:collapse;width:100%;min-width:760px;}
  table.sched th{font-size:11px;text-transform:uppercase;letter-spacing:0.04em;color:var(--muted);font-weight:700;padding:8px 6px;border-bottom:1px solid var(--line);text-align:left;}
  table.sched td.timecol{font-family:'IBM Plex Mono',monospace;font-size:10.5px;color:var(--muted);white-space:nowrap;padding:10px 8px 10px 0;vertical-align:top;}
  table.sched td.cell{border:1px solid var(--line);padding:0;vertical-align:top;width:19%;}
  .cell-inner{padding:8px 9px;min-height:58px;font-size:12px;}
  .cell-inner.event{border-left:3px solid #C9A24A; background:var(--amber-soft);}
  .cell-inner.task{border-left:3px solid var(--accent);}
  .cell-inner.empty{border-left:3px solid var(--line-strong); background:#FBFBF9; color:var(--muted); font-size:11px;}
  .cell-task{font-weight:600;line-height:1.4; outline:none; border-radius:5px; padding:1px 3px; margin:-1px -3px; cursor:text;}
  .cell-task:hover{background:rgba(15,102,87,0.06);}
  .cell-task:focus{background:var(--surface);box-shadow:0 0 0 2px var(--accent-soft);}
  .cell-reason{font-size:10.3px;color:var(--muted);margin-top:4px;line-height:1.4;}
  .cell-badge{display:inline-block;font-size:9px;font-weight:700;padding:1px 6px;border-radius:999px;margin-top:5px;font-family:'IBM Plex Mono',monospace;background:var(--line); color:var(--muted);}
  .cell-badge.pri-긴급, .cell-badge.pri-상{background:var(--warn-soft);color:var(--warn);}
  .cell-badge.pri-중{background:var(--accent-soft);color:var(--accent-dark);}
  .cell-badge.edited{background:#EDE4F5;color:#6A4C93;margin-left:4px;}
  tr.lunch td.cell .cell-inner{background:#F1F1EC;color:var(--muted);font-size:11px;text-align:center;border-left:3px solid var(--line-strong);}
  .overflow-box{margin-top:14px;padding:12px 14px;border-radius:10px;border:1px solid var(--line);background:#FBFBF9;font-size:12px;}
  .overflow-box h4{margin:0 0 6px;font-size:11.5px;color:var(--muted);text-transform:uppercase;letter-spacing:0.04em;}
  .overflow-box ul{margin:0;padding-left:18px;}
  .mail-box{margin-top:10px;}
  .mail-box summary{cursor:pointer;font-size:11.5px;color:var(--muted);font-weight:600;}
  .mail-box .mail-item{font-size:11.5px;padding:6px 0;border-bottom:1px dashed var(--line);}
  .hint{font-size:11.8px;color:var(--muted);margin-top:4px;}
  .ai-note{font-size:12px;line-height:1.6;background:var(--accent-soft);border:1px solid #CFE0FF;color:var(--accent-dark);border-radius:10px;padding:10px 12px;margin-bottom:12px;}
  .error{color:var(--warn);font-size:12px;margin-top:8px;}
  .empty-msg{font-size:13px;color:var(--muted);padding:24px 0;text-align:center;}
  .legend{display:flex;gap:14px;flex-wrap:wrap;margin-bottom:10px;}
  .legend span{font-size:10.6px;color:var(--muted);display:flex;align-items:center;gap:5px;}
  .legend .sw{width:9px;height:9px;border-radius:3px;display:inline-block;}
  .legend .sw.event{background:#C9A24A;}
  .legend .sw.task{background:var(--accent);}
  .legend .sw.empty{background:var(--line-strong);}
  .role-toggle{display:flex;align-items:center;gap:8px;margin-bottom:12px;flex-wrap:wrap;}
  .role-toggle .rt-label{font-size:11.5px;color:var(--muted);font-weight:600;}
  .role-chip{font-size:12px;font-weight:600;padding:6px 12px;border-radius:999px;border:1px solid var(--line-strong);background:var(--surface);cursor:pointer;}
  .role-chip.active{background:var(--ink);color:white;border-color:var(--ink);}
  .role-chip.active.mentor{background:#6A4C93;border-color:#6A4C93;}
  .cell-badge.edited.employee{background:var(--accent-soft);color:var(--accent-dark);}
  .cell-badge.edited.mentor{background:#EDE4F5;color:#6A4C93;}
</style>
</head>
<body>

<header>
  <div><span class="mark">H</span><h1>후임자 온보딩 일정 설계</h1></div>
  <p>회의록 · 캘린더 · 메일 문서의 실제 일정과 업무 자료를 합쳐서 <b>1주차는 최대한 빼곡하게</b>, <b>2주차는 실제 날짜에 있는 일정만</b> 반영한 일과표(월~금)를 만들어줍니다. 생성 후 칸을 클릭해 바로 수정할 수 있어요.</p>
</header>

<div class="layout">

  <!-- LEFT -->
  <div class="panel">
    <div class="panel-head"><h2>문서</h2></div>
    <div class="panel-body">

      <div id="aiDocNote" class="ai-note" style="display:none;"></div>

      <div id="manualUpload">
        <div class="source-block">
          <div class="source-label">회의록 <span class="req">(여러 개, docx)</span></div>
          <div class="dropzone" data-target="meeting">
            <div class="dz-title">파일 업로드</div>
            <div class="dz-sub">.docx (여러 개 선택 가능)</div>
          </div>
          <input type="file" class="fileInput" data-target="meeting" accept=".docx" multiple>
          <div class="file-list" id="fileList-meeting"></div>
        </div>

        <div class="source-block">
          <div class="source-label">캘린더 문서 <span class="req">(06_인수인계_캘린더.docx)</span></div>
          <div class="dropzone" data-target="calendar">
            <div class="dz-title">파일 업로드</div>
            <div class="dz-sub">.docx</div>
          </div>
          <input type="file" class="fileInput" data-target="calendar" accept=".docx">
          <div class="file-list" id="fileList-calendar"></div>
        </div>

        <div class="source-block">
          <div class="source-label">메일 <span class="req">(여러 개, 참고용 표시)</span></div>
          <div class="dropzone" data-target="email">
            <div class="dz-title">파일 업로드</div>
            <div class="dz-sub">.docx (여러 개 선택 가능)</div>
          </div>
          <input type="file" class="fileInput" data-target="email" accept=".docx" multiple>
          <div class="file-list" id="fileList-email"></div>
        </div>
      </div>

      <div class="row-btns">
        <button class="btn btn-primary btn-block" id="genBtn">일정표 생성</button>
      </div>
      <p class="hint">업무일정·프로젝트 현황·자산·연락망 내용은 도구 안에 이미 반영되어 있어서 따로 업로드하지 않아도 됩니다.</p>
    </div>
  </div>

  <!-- RIGHT -->
  <div class="panel">
    <div class="panel-head"><h2>1주차 · 2주차 일과표</h2></div>
    <div class="panel-body">

      <div class="legend">
        <span><span class="sw event"></span>회의록/캘린더의 고정 시간 일정</span>
        <span><span class="sw task"></span>문서에서 온 업무/개요 항목</span>
        <span><span class="sw empty"></span>여유 시간</span>
      </div>

      <div class="role-toggle">
        <span class="rt-label">지금 누가 수정하나요?</span>
        <div class="role-chip active" id="role-employee" data-role="employee">신입</div>
        <div class="role-chip" id="role-mentor" data-role="mentor">선임자</div>
      </div>
      <p class="hint" style="margin-top:-6px;">신입 또는 선임자 버튼을 누른 뒤 칸을 클릭해 수정하면, 누가 고쳤는지 칸 아래 배지로 표시됩니다.</p>

      <div id="weekTabs" class="week-tabs"></div>

      <div class="row-btns" style="margin-top:0;margin-bottom:12px;">
        <button class="btn btn-ghost" id="resetBtn" style="display:none;">이 주 수정 초기화</button>
      </div>

      <div id="gridArea">
        <div class="empty-msg">"일정표 생성"을 눌러보세요.</div>
      </div>

      <div id="overflowArea"></div>
      <div id="mailArea"></div>
      <div id="laterArea"></div>

      <div id="errorBox" class="error" style="display:none;"></div>
    </div>
  </div>

</div>

<script>
const SLOTS = ["09:00-10:00","10:00-12:00","13:00-15:00","15:00-17:00","17:00-18:00"];
const SLOT_RANGES = [[540,600],[600,720],[780,900],[900,1020],[1020,1080]]; // minutes from midnight

// ===== Streamlit에서 AI가 분석한 문서 데이터 (Python이 값을 주입) =====
const AI_DOC_MODE = false;
const AI_START_DATE = "";
const AI_TIMED_EVENTS = [];
const AI_ALLDAY_TASKS = [];
const AI_MAIL_REFS = [];

let files = { meeting:[], calendar:[], email:[] };
let fileIdCounter = 0;

let originalWeeks = null; // parsed, immutable
let workingWeeks = null;  // editable copy
let weekOrder = [];
let currentWeekKey = null;
let currentRole = 'employee'; // 'employee' | 'mentor'

const el = id => document.getElementById(id);
const escapeHtml = s => (s||'').replace(/[&<>"']/g, c => ({'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;',"'":'&#39;'}[c]));

function showError(msg){ el('errorBox').style.display='block'; el('errorBox').textContent = msg; }
function clearError(){ el('errorBox').style.display='none'; }

if(AI_DOC_MODE){
  el('manualUpload').style.display = 'none';
  const note = el('aiDocNote');
  note.style.display = 'block';
  note.innerHTML = `🤖 <b>AI가 분석한 문서가 반영됩니다.</b><br>시간 일정 ${AI_TIMED_EVENTS.length}건 · 마감/조치 ${AI_ALLDAY_TASKS.length}건 · 메일 ${AI_MAIL_REFS.length}건`
    + (AI_START_DATE ? `<br>인수 시작일(1주차 기준): <b>${escapeHtml(AI_START_DATE)}</b>` : '');
}

/* ---------------- 파일 업로드 UI (AI 미사용 시) ---------------- */

function renderFileList(target){
  const wrap = el('fileList-'+target);
  wrap.innerHTML = '';
  files[target].forEach(d=>{
    const row = document.createElement('div');
    row.className = 'file-item';
    const statusText = d.status === 'loading' ? '읽는 중…' : d.status === 'ok' ? '완료' : '실패';
    const statusClass = d.status === 'ok' ? 'ok' : d.status === 'err' ? 'err' : '';
    row.innerHTML = `<span class="fname">${escapeHtml(d.name)}</span><span class="fstatus ${statusClass}">${statusText}</span><span class="fremove" data-id="${d.id}">✕</span>`;
    row.querySelector('.fremove').onclick = () => {
      files[target] = files[target].filter(x => x.id !== d.id);
      renderFileList(target);
    };
    wrap.appendChild(row);
  });
}

async function handleFiles(target, fileListRaw, single){
  const arr = Array.from(fileListRaw);
  if(single){ files[target] = []; }
  for(const file of arr){
    const doc = { id: ++fileIdCounter, name: file.name, status: 'loading', raw:null };
    files[target].push(doc);
    renderFileList(target);
    try{
      const buf = await file.arrayBuffer();
      const res = await mammoth.extractRawText({ arrayBuffer: buf });
      doc.raw = res.value;
      doc.status = 'ok';
    }catch(err){
      doc.status = 'err';
      showError(`${file.name} 처리 실패: ${err.message}`);
    }
    renderFileList(target);
  }
}

document.querySelectorAll('.dropzone').forEach(dz=>{
  const target = dz.dataset.target;
  const input = document.querySelector(`.fileInput[data-target="${target}"]`);
  const single = !input.multiple;
  dz.addEventListener('click', ()=> input.click());
  input.addEventListener('change', e=>{ handleFiles(target, e.target.files, single); input.value=''; });
  ['dragenter','dragover'].forEach(evt=> dz.addEventListener(evt, e=>{ e.preventDefault(); dz.classList.add('drag'); }));
  ['dragleave','drop'].forEach(evt=> dz.addEventListener(evt, e=>{ e.preventDefault(); dz.classList.remove('drag'); }));
  dz.addEventListener('drop', e=>{ if(e.dataTransfer.files.length) handleFiles(target, e.dataTransfer.files, single); });
});

document.querySelectorAll('.role-chip').forEach(chip=>{
  chip.addEventListener('click', ()=>{
    currentRole = chip.dataset.role;
    document.querySelectorAll('.role-chip').forEach(c=>c.classList.remove('active','mentor'));
    chip.classList.add('active');
    if(currentRole === 'mentor') chip.classList.add('mentor');
  });
});

/* ---------------- 날짜/시간 유틸 ---------------- */

// 로컬 날짜를 YYYY-MM-DD 로 (toISOString 은 UTC 라서 한국 시간에서 하루 밀리는 문제가 있음)
function ymd(d){
  return `${d.getFullYear()}-${String(d.getMonth()+1).padStart(2,'0')}-${String(d.getDate()).padStart(2,'0')}`;
}

function normalizeDate(v){
  if(v instanceof Date && !isNaN(v)) return ymd(v);
  if(typeof v === 'string'){
    let m = v.match(/(\d{4})-(\d{1,2})-(\d{1,2})/);
    if(m) return `${m[1]}-${String(m[2]).padStart(2,'0')}-${String(m[3]).padStart(2,'0')}`;
    m = v.match(/^(\d{1,2})\/(\d{1,2})$/);
    if(m){
      const y = AI_START_DATE ? AI_START_DATE.slice(0,4) : String(new Date().getFullYear());
      return `${y}-${String(m[1]).padStart(2,'0')}-${String(m[2]).padStart(2,'0')}`;
    }
  }
  return null;
}

function timeToMinutes(hhmm){
  const m = (hhmm||'').match(/(\d{1,2}):(\d{2})/);
  if(!m) return null;
  return (+m[1])*60 + (+m[2]);
}

function slotIndexForTime(hhmm){
  const min = timeToMinutes(hhmm);
  if(min===null) return 0;
  for(let i=0;i<SLOT_RANGES.length;i++){
    if(min >= SLOT_RANGES[i][0] && min < SLOT_RANGES[i][1]) return i;
  }
  if(min < SLOT_RANGES[0][0]) return 0;
  return SLOT_RANGES.length-1;
}

function weekKeyOf(dateStr){
  const d = new Date(dateStr+'T00:00:00');
  const day = d.getDay(); // 0=Sun
  const diffToMon = (day===0? -6 : 1-day);
  const mon = new Date(d); mon.setDate(d.getDate()+diffToMon);
  return ymd(mon);
}
function addDays(dateStr, n){
  const d = new Date(dateStr+'T00:00:00');
  d.setDate(d.getDate()+n);
  return ymd(d);
}

/* ---------------- 내장 데이터 (Excel 업로드 시 Python이 교체) ---------------- */

const EMBEDDED_SCHEDULE_TASKS = [
  { date:"2026-09-01", title:"최종 점검보고서 제출", priority:"긴급", nextAction:"누락 사진 3장 확인 → 팀장 검토 → 고객사 제출", note:"오전 중 제출 권장", source:"업무일정 엑셀" },
  { date:"2026-09-02", title:"감지기 교체 견적 회신", priority:"상", nextAction:"한빛전기 단가 회신 확인 후 견적서 최종 작성", note:"협력사 회신 대기", source:"업무일정 엑셀" },
  { date:"2026-09-03", title:"고객사 정기회의 참석", priority:"상", nextAction:"회의 전 미조치 2건 및 사진자료 정리", note:"회의 14:00", source:"업무일정 엑셀" },
  { date:"2026-09-04", title:"주간 미완료 업무 점검", priority:"중", nextAction:"미완료 보고서/견적/고객 요청 목록 업데이트", note:"매주 금요일 반복", source:"업무일정 엑셀" },
  { date:"2026-09-07", title:"공용드라이브 권한 이관 요청", priority:"상", nextAction:"팀장 승인 후 이서연 계정에 편집권한 부여 요청", note:"외부공유 권한 제외", source:"업무일정 엑셀" },
  { date:"2026-09-08", title:"월간 점검 일정 확정", priority:"중", nextAction:"현장팀 일정 취합 후 9월 점검표 확정", note:"현장팀 3명 일정 확인 필요", source:"업무일정 엑셀" },
  { date:"2026-09-10", title:"D물류센터 사전자료 요청", priority:"중", nextAction:"도면/설비목록/이전 점검결과 요청 메일 발송", note:"신규 인계 후 첫 신규 현장", source:"업무일정 엑셀" },
];

const EMBEDDED_PROJECT_DEADLINE_TASKS = [
  { date:"2026-09-01", title:"[마감] A동 소방시설 정기점검", priority:"상", nextAction:"누락 사진 3장 확인 후 최종 제출", note:"사진 누락 시 제출 지연 가능", source:"프로젝트 진행현황 엑셀" },
  { date:"2026-09-02", title:"[마감] B공장 감지기 교체", priority:"상", nextAction:"협력사 단가 반영 후 견적 회신", note:"단가 지연 시 고객 회신 지연", source:"프로젝트 진행현황 엑셀" },
  { date:"2026-09-03", title:"[마감] C센터 종합정밀점검", priority:"상", nextAction:"9/3 회의에서 미조치 2건 일정 확정", note:"미조치 2건 일정 미확정", source:"프로젝트 진행현황 엑셀" },
  { date:"2026-09-10", title:"[마감] D물류센터 신규점검", priority:"상", nextAction:"사전자료 요청 및 현장 일정 협의", note:"자료 미수신 시 현장 준비 지연", source:"프로젝트 진행현황 엑셀" },
];

const EMBEDDED_ASSET_DEADLINE_TASKS = [
  { date:"2026-09-07", title:"[마감] 공용드라이브 Z: 인계", priority:"상", nextAction:"팀장 승인 후 IT 요청", note:"외부 공유 권한은 부여하지 않음", source:"계정·권한·자산 엑셀" },
  { date:"2026-09-04", title:"[마감] 고객요청 관리 엑셀 인계", priority:"상", nextAction:"최신 파일 경로 전달", note:"중복본 사용 금지", source:"계정·권한·자산 엑셀" },
  { date:"2026-09-07", title:"[마감] A동 현장 폴더 인계", priority:"상", nextAction:"공용드라이브 권한과 함께 이관", note:"보고서 최종본 폴더 확인", source:"계정·권한·자산 엑셀" },
  { date:"2026-09-04", title:"[마감] 법인 태블릿 2번 인계", priority:"상", nextAction:"자산대장 서명 후 인계", note:"충전기 포함", source:"계정·권한·자산 엑셀" },
];

const EMBEDDED_PROJECT_OVERVIEW = [
  { title:"[개요] A동 소방시설 정기점검 현황 파악", detail:"보고서 최종화 · 진행률 80% · 다음액션: 누락 사진 3장 확인 후 최종 제출", badge:"개요", source:"프로젝트 진행현황 엑셀" },
  { title:"[개요] B공장 감지기 교체 현황 파악", detail:"견적 작성 · 진행률 60% · 다음액션: 협력사 단가 반영 후 견적 회신", badge:"개요", source:"프로젝트 진행현황 엑셀" },
  { title:"[개요] C센터 종합정밀점검 현황 파악", detail:"후속조치 협의 · 진행률 30% · 다음액션: 9/3 회의에서 미조치 2건 일정 확정", badge:"개요", source:"프로젝트 진행현황 엑셀" },
  { title:"[개요] D물류센터 신규점검 현황 파악", detail:"사전준비 · 진행률 10% · 다음액션: 사전자료 요청 및 현장 일정 협의", badge:"개요", source:"프로젝트 진행현황 엑셀" },
];

const EMBEDDED_ASSET_OVERVIEW = [
  { title:"[개요] 공용드라이브 Z: 인계 상태 확인", detail:"상태: 인계예정 · 인계방법: 팀장 승인 후 IT 요청 · 주의사항: 외부 공유 권한은 부여하지 않음", badge:"개요", source:"계정·권한·자산 엑셀" },
  { title:"[개요] 고객요청 관리 엑셀 인계 상태 확인", detail:"상태: 미완료 · 인계방법: 최신 파일 경로 전달 · 주의사항: 중복본 사용 금지", badge:"개요", source:"계정·권한·자산 엑셀" },
  { title:"[개요] 사내 IT요청 포털 인계 상태 확인", detail:"상태: 완료 · 인계방법: 개인 계정 직접 로그인 · 주의사항: 비밀번호 공유 금지", badge:"개요", source:"계정·권한·자산 엑셀" },
  { title:"[개요] A동 현장 폴더 인계 상태 확인", detail:"상태: 인계예정 · 인계방법: 공용드라이브 권한과 함께 이관 · 주의사항: 보고서 최종본 폴더 확인", badge:"개요", source:"계정·권한·자산 엑셀" },
  { title:"[개요] 법인 태블릿 2번 인계 상태 확인", detail:"상태: 미완료 · 인계방법: 자산대장 서명 후 인계 · 주의사항: 충전기 포함", badge:"개요", source:"계정·권한·자산 엑셀" },
];

const EMBEDDED_CONTACTS = [
  { title:"박현우(부장, 세림관리) 컨택포인트 파악", detail:"관련 업무: A동 점검보고서 · 유의사항: 보고서 전달 전 전화로 먼저 안내. 오전 11시 이전 연락 선호.", badge:"소개", source:"담당자 연락망 엑셀" },
  { title:"최은지(대리, B공장 시설팀) 컨택포인트 파악", detail:"관련 업무: B공장 견적/교체 일정 · 유의사항: 메일 제목에 [B공장] 표기. 견적 수정사항은 표로 정리.", badge:"소개", source:"담당자 연락망 엑셀" },
  { title:"조민석(과장, 한빛전기) 컨택포인트 파악", detail:"관련 업무: 감지기 단가/납기 · 유의사항: 급한 건 전화, 일반 단가 문의는 문자 가능.", badge:"소개", source:"담당자 연락망 엑셀" },
  { title:"오세훈(과장, C센터 시설팀) 컨택포인트 파악", detail:"관련 업무: C센터 후속조치 · 유의사항: 회의자료는 전날 17시까지 공유.", badge:"소개", source:"담당자 연락망 엑셀" },
  { title:"박지훈(대리, 현장점검팀) 컨택포인트 파악", detail:"관련 업무: 현장 사진/점검결과 · 유의사항: 사진 누락 확인은 박지훈 대리에게 요청.", badge:"소개", source:"담당자 연락망 엑셀" },
  { title:"한유리(사원, 영업팀) 컨택포인트 파악", detail:"관련 업무: 견적/고객 요청 · 유의사항: 견적 금액 변경 시 반드시 공유.", badge:"소개", source:"담당자 연락망 엑셀" },
];

/* ---------------- DOCX(텍스트) 파싱 — AI 미사용 시 예비 ---------------- */

function splitLines(rawText){
  return rawText.split('\n').map(s=>s.trim()).filter(Boolean);
}

function parseCalendarDocx(rawText){
  const lines = splitLines(rawText);
  let start = -1;
  for(let i=0;i<lines.length-3;i++){
    if(lines[i]==='날짜' && lines[i+1]==='시간' && lines[i+2]==='일정' && lines[i+3]==='내용'){ start = i+4; break; }
  }
  const events = [];
  if(start<0) return events;
  for(let i=start;i+3<lines.length;i+=4){
    const [dateStr, timeStr, title, detail] = lines.slice(i,i+4);
    const date = normalizeDate(dateStr);
    if(!date) break;
    const tm = timeStr.match(/(\d{1,2}:\d{2})\s*-\s*(\d{1,2}:\d{2})/);
    events.push({ date, start: tm ? tm[1] : null, end: tm ? tm[2] : null, allDay: !tm, title, detail, source: '캘린더 문서' });
  }
  return events;
}

function parseMeetingDocx(rawText, filename){
  const lines = splitLines(rawText);
  if(!lines.length) return null;
  const title = lines[0];
  let dtLine = null, place='', attendees='';
  for(let i=0;i<lines.length;i++){
    if(lines[i]==='일시' && lines[i+1]) dtLine = lines[i+1];
    if(lines[i]==='장소' && lines[i+1]) place = lines[i+1];
    if(lines[i]==='참석자' && lines[i+1]) attendees = lines[i+1];
  }
  if(!dtLine) return null;
  const m = dtLine.match(/(\d{4}-\d{2}-\d{2})\s+(\d{1,2}:\d{2})\s*-\s*(\d{1,2}:\d{2})/);
  if(!m) return null;
  return { date: normalizeDate(m[1]), start: m[2], end: m[3], allDay:false, title,
           detail: `참석: ${attendees}${place? ' · 장소: '+place : ''}`, source: '회의록 (' + filename + ')' };
}

function parseEmailDocx(rawText, filename){
  const lines = splitLines(rawText);
  let subject='', to='', dt='';
  for(let i=0;i<lines.length;i++){
    if(lines[i]==='제목' && lines[i+1]) subject = lines[i+1];
    if(lines[i]==='받는사람' && lines[i+1]) to = lines[i+1];
    if(lines[i]==='일시' && lines[i+1]) dt = lines[i+1];
  }
  const date = normalizeDate(dt);
  if(!subject) return null;
  return { date, subject, to, raw: dt, source: filename };
}

/* ---------------- 일정표 생성 ---------------- */

function buildSchedule(){
  const scheduleTasks = [ ...EMBEDDED_SCHEDULE_TASKS, ...EMBEDDED_PROJECT_DEADLINE_TASKS, ...EMBEDDED_ASSET_DEADLINE_TASKS ];
  const timedEvents = [];
  const mailRefs = [];

  // (1) AI가 분석한 문서 데이터
  AI_TIMED_EVENTS.forEach(ev => timedEvents.push({ date:ev.date, start:ev.start, end:ev.end, allDay:false, title:ev.title, detail:ev.detail, source:ev.source }));
  AI_ALLDAY_TASKS.forEach(t => scheduleTasks.push({ date:t.date, title:t.title, priority:t.priority||'중', nextAction:t.detail||'', note:'', source:t.source }));
  AI_MAIL_REFS.forEach(m => mailRefs.push({ date:m.date, subject:m.subject, to:m.to, raw:[m.date, m.time].filter(Boolean).join(' '), summary:m.summary||'', source:m.source }));

  // (2) 이 화면에서 직접 올린 문서 (AI 미사용 시)
  files.calendar.filter(f=>f.status==='ok').forEach(f => {
    parseCalendarDocx(f.raw).forEach(ev => {
      if(ev.allDay){
        scheduleTasks.push({ date:ev.date, title:ev.title, priority:'상', nextAction: ev.detail, note:'', source: ev.source });
      } else {
        timedEvents.push(ev);
      }
    });
  });
  files.meeting.filter(f=>f.status==='ok').forEach(f => { const ev = parseMeetingDocx(f.raw, f.name); if(ev) timedEvents.push(ev); });
  files.email.filter(f=>f.status==='ok').forEach(f => { const m = parseEmailDocx(f.raw, f.name); if(m) mailRefs.push(m); });

  if(scheduleTasks.length===0 && timedEvents.length===0){
    return null;
  }

  const priorityRank = { '긴급':0, '상':1, '중':2, '하':3, '':2 };
  scheduleTasks.sort((a,b) => (priorityRank[a.priority]??2) - (priorityRank[b.priority]??2));

  // 같은 날짜의 거의 같은 업무(한쪽 이름이 다른 쪽에 포함)는 한 번만 표시
  const normTitle = s => (s||'').replace(/\[[^\]]*\]/g,'').replace(/[\s()·\-:]/g,'');
  const dedupedTasks = [];
  scheduleTasks.forEach(t => {
    const n = normTitle(t.title);
    const dup = n && dedupedTasks.find(x => {
      if(x.date !== t.date) return false;
      const xn = normTitle(x.title);
      return xn && (xn.includes(n) || n.includes(xn));
    });
    if(!dup) dedupedTasks.push(t);
  });

  // 같은 날짜+같은 시작시간 일정은 중복으로 보고 하나만 사용
  const seenEventKeys = new Set();
  const dedupedEvents = [];
  timedEvents.forEach(ev => {
    const key = ev.date + '|' + (ev.start || 'allday');
    if(seenEventKeys.has(key)) return;
    seenEventKeys.add(key);
    dedupedEvents.push(ev);
  });

  const tasksByDate = {};
  dedupedTasks.forEach(t => { if(t.date) (tasksByDate[t.date] = tasksByDate[t.date]||[]).push(t); });
  const eventsByDate = {};
  dedupedEvents.forEach(e => { if(e.date) (eventsByDate[e.date] = eventsByDate[e.date]||[]).push(e); });
  const mailByDate = {};
  mailRefs.forEach(m => { if(m.date) (mailByDate[m.date] = mailByDate[m.date]||[]).push(m); });

  const allDates = Array.from(new Set([...Object.keys(tasksByDate), ...Object.keys(eventsByDate)])).sort();

  // 1주차 기준: AI가 문서에서 찾은 '인수 시작일'이 있으면 그 주, 없으면 가장 이른 날짜의 주
  let usedWeekKeys;
  if(AI_START_DATE){
    const w1 = weekKeyOf(AI_START_DATE);
    usedWeekKeys = [w1, weekKeyOf(addDays(w1, 7))];
  } else {
    usedWeekKeys = Array.from(new Set(allDates.map(weekKeyOf))).sort().slice(0, 2);
  }
  const week1Key = usedWeekKeys[0];
  const week2Key = usedWeekKeys[1];
  const laterDates = allDates.filter(d => !usedWeekKeys.includes(weekKeyOf(d)));

  const DAY_NAMES = ['월','화','수','목','금'];
  const weeks = {};

  function emptySlots(){
    return SLOTS.map(s => ({ slot:s, type:'empty', task:'', detail:'', badge:'', source:'', edited:false, editedBy:null }));
  }
  function placeIntoDay(dayObj, item){
    const freeIdx = dayObj.slots.findIndex(s => s.type==='empty');
    if(freeIdx>=0){
      dayObj.slots[freeIdx] = { slot:SLOTS[freeIdx], type:'task', task:item.title, detail:item.detail||'', badge:item.badge||'', source:item.source||'', edited:false, editedBy:null };
      return true;
    }
    dayObj.overflow.push(`${item.title} — ${item.source||''}`);
    return false;
  }
  function dayNameOf(date){
    const idx = new Date(date+'T00:00:00').getDay() - 1;
    return (idx < 0 || idx > 4) ? null : DAY_NAMES[idx];
  }

  // ---------- 2주차: 실제 날짜에 있는 항목만 반영 ----------
  if(week2Key){
    weeks[week2Key] = { label: '2주차', days:{}, mail:{} };
    allDates.forEach(date => {
      if(weekKeyOf(date) !== week2Key) return;
      const dayName = dayNameOf(date);
      if(!dayName) return;
      if(!weeks[week2Key].days[dayName]) weeks[week2Key].days[dayName] = { slots: emptySlots(), overflow: [] };
      const dayObj = weeks[week2Key].days[dayName];

      (eventsByDate[date]||[]).forEach(ev => {
        const idx = slotIndexForTime(ev.start);
        if(dayObj.slots[idx].type==='empty'){
          dayObj.slots[idx] = { slot:SLOTS[idx], type:'event', task:`${ev.start}-${ev.end} ${ev.title}`, detail: ev.detail||'', badge:'', source: ev.source, edited:false, editedBy:null };
        } else {
          dayObj.overflow.push(`(추가 일정) ${ev.start}-${ev.end} ${ev.title} — ${ev.source}`);
        }
      });
      (tasksByDate[date]||[]).forEach(t => {
        placeIntoDay(dayObj, { title:t.title, detail: t.nextAction ? `다음 조치: ${t.nextAction}` : (t.note||''), badge:t.priority, source:t.source });
      });
    });
    Object.keys(mailByDate).forEach(date => {
      if(weekKeyOf(date) !== week2Key) return;
      const dayName = dayNameOf(date);
      if(dayName) weeks[week2Key].mail[dayName] = (weeks[week2Key].mail[dayName]||[]).concat(mailByDate[date]);
    });
  }

  // ---------- 1주차: 문서 내용을 최대한 빼곡하게 채운 초안 ----------
  if(week1Key){
    weeks[week1Key] = { label: '1주차', days:{}, mail:{} };
    DAY_NAMES.forEach(dn => { weeks[week1Key].days[dn] = { slots: emptySlots(), overflow: [] }; });

    // 1) 시간이 정해진 회의/캘린더 일정은 그대로 고정 배치
    allDates.forEach(date => {
      if(weekKeyOf(date) !== week1Key) return;
      const dayName = dayNameOf(date);
      if(!dayName) return;
      const dayObj = weeks[week1Key].days[dayName];
      (eventsByDate[date]||[]).forEach(ev => {
        const idx = slotIndexForTime(ev.start);
        if(dayObj.slots[idx].type==='empty'){
          dayObj.slots[idx] = { slot:SLOTS[idx], type:'event', task:`${ev.start}-${ev.end} ${ev.title}`, detail: ev.detail||'', badge:'', source: ev.source, edited:false, editedBy:null };
        }
      });
    });
    Object.keys(mailByDate).forEach(date => {
      if(weekKeyOf(date) !== week1Key) return;
      const dayName = dayNameOf(date);
      if(dayName) weeks[week1Key].mail[dayName] = (weeks[week1Key].mail[dayName]||[]).concat(mailByDate[date]);
    });

    // 2) 1주차의 실제 마감 업무를 먼저, 그다음 학습 풀로 남은 칸을 채운다
    const week1Tasks = [];
    allDates.forEach(date => {
      if(weekKeyOf(date) !== week1Key) return;
      (tasksByDate[date]||[]).forEach(t => week1Tasks.push({ date, t }));
    });
    const placedTitles = new Set();
    week1Tasks.forEach(({date, t}) => {
      const dayName = dayNameOf(date);
      if(!dayName) return;
      placeIntoDay(weeks[week1Key].days[dayName], { title:t.title, detail: t.nextAction ? `다음 조치: ${t.nextAction}` : (t.note||''), badge:t.priority, source:t.source });
      placedTitles.add(t.title);
    });

    const pool = [];
    pool.push(...EMBEDDED_PROJECT_OVERVIEW);
    EMBEDDED_SCHEDULE_TASKS.forEach(t => {
      if(placedTitles.has(t.title)) return;
      pool.push({ title: `실습: ${t.title}`, detail: t.nextAction ? `다음 조치: ${t.nextAction}` : (t.note||''), badge: '실습', source: t.source });
    });
    pool.push(...EMBEDDED_ASSET_OVERVIEW);
    pool.push(...EMBEDDED_CONTACTS);
    mailRefs.forEach(m => pool.push({
      title: `메일 확인: ${m.subject}`,
      detail: m.summary ? m.summary : `수신: ${m.to}${m.raw ? ' · '+m.raw : ''}`,
      badge: '메일', source: `메일 문서(${m.source})`,
    }));

    let pi = 0;
    DAY_NAMES.forEach(dn => {
      const dayObj = weeks[week1Key].days[dn];
      for(let si=0; si<SLOTS.length; si++){
        if(dayObj.slots[si].type !== 'empty') continue;
        if(pi >= pool.length) break;
        placeIntoDay(dayObj, pool[pi]);
        pi++;
      }
    });
    if(pi < pool.length){
      weeks[week1Key].days['금'].overflow.push(...pool.slice(pi).map(p => `${p.title} — ${p.source||''}`));
    }
  }

  return { weeks, laterDates, weekOrder: [week1Key, week2Key].filter(Boolean) };
}

function deepCopy(obj){ return JSON.parse(JSON.stringify(obj)); }

function generate(){
  clearError();
  const result = buildSchedule();
  if(!result){
    showError('일정표를 만들지 못했습니다. 업무자료나 문서를 먼저 올려주세요.');
    return;
  }
  originalWeeks = result.weeks;
  workingWeeks = deepCopy(result.weeks);
  weekOrder = result.weekOrder;
  currentWeekKey = weekOrder[0];
  renderWeekTabs();
  renderWeek(currentWeekKey);
  renderLaterNote(result.laterDates);
}

function renderLaterNote(laterDates){
  const host = el('laterArea');
  if(!laterDates || !laterDates.length){ host.innerHTML=''; return; }
  host.innerHTML = `<div class="overflow-box"><h4>1~2주차 범위 밖 일정 (${laterDates.length}일, 표에는 반영되지 않음)</h4><div class="hint">${laterDates.map(escapeHtml).join(', ')}</div></div>`;
}

function renderWeekTabs(){
  const wrap = el('weekTabs');
  wrap.innerHTML = '';
  weekOrder.forEach(wk => {
    const tab = document.createElement('div');
    tab.className = 'week-tab' + (wk===currentWeekKey ? ' active':'');
    const mon = new Date(wk+'T00:00:00');
    const fri = new Date(mon); fri.setDate(mon.getDate()+4);
    tab.innerHTML = `${workingWeeks[wk].label} <span style="font-weight:400;color:var(--muted);font-size:11px;">${mon.getMonth()+1}/${mon.getDate()}~${fri.getMonth()+1}/${fri.getDate()}</span>`;
    tab.onclick = () => { currentWeekKey = wk; renderWeekTabs(); renderWeek(wk); };
    wrap.appendChild(tab);
  });
}

const DAY_NAMES_FULL = { '월':'월요일', '화':'화요일', '수':'수요일', '목':'목요일', '금':'금요일' };

function renderWeek(wk){
  const weekData = workingWeeks[wk];
  const dayKeys = ['월','화','수','목','금'];
  el('resetBtn').style.display = 'inline-block';

  let html = '<div class="grid-wrap"><table class="sched"><thead><tr><th></th>';
  dayKeys.forEach(dk => html += `<th>${DAY_NAMES_FULL[dk]}</th>`);
  html += '</tr></thead><tbody>';

  SLOTS.forEach((slot, si) => {
    html += `<tr><td class="timecol">${slot}</td>`;
    dayKeys.forEach(dk => {
      const dayEntry = weekData.days[dk];
      const cellData = dayEntry ? dayEntry.slots[si] : { type:'empty', task:'', detail:'', badge:'', source:'', edited:false };
      const cls = cellData.type;
      html += `<td class="cell"><div class="cell-inner ${cls}" data-day="${dk}" data-slotidx="${si}">`;
      if(cellData.type==='empty'){
        html += `<div>여유 시간</div>`;
      } else {
        html += `<div class="cell-task" contenteditable="true" spellcheck="false">${escapeHtml(cellData.task)}</div>`;
        if(cellData.detail) html += `<div class="cell-reason">${escapeHtml(cellData.detail)}</div>`;
        if(cellData.badge) html += `<span class="cell-badge pri-${escapeHtml(cellData.badge)}">${escapeHtml(cellData.badge)}</span>`;
        if(cellData.editedBy) html += `<span class="cell-badge edited ${cellData.editedBy}">✎ ${cellData.editedBy==='mentor'?'선임자':'신입'} 수정</span>`;
        if(cellData.source) html += `<div class="cell-reason" style="opacity:0.7;">출처: ${escapeHtml(cellData.source)}</div>`;
      }
      html += `</div></td>`;
    });
    html += '</tr>';
    if(slot === '10:00-12:00'){
      html += `<tr class="lunch"><td class="timecol">12:00-13:00</td>`;
      dayKeys.forEach(()=> html += `<td class="cell"><div class="cell-inner">점심시간</div></td>`);
      html += `</tr>`;
    }
  });

  html += '</tbody></table></div>';
  el('gridArea').innerHTML = html;

  const overflowLines = [];
  dayKeys.forEach(dk => {
    const dayEntry = weekData.days[dk];
    if(dayEntry && dayEntry.overflow.length){
      dayEntry.overflow.forEach(line => overflowLines.push(`${DAY_NAMES_FULL[dk]}: ${line}`));
    }
  });
  el('overflowArea').innerHTML = overflowLines.length
    ? `<div class="overflow-box"><h4>이번 주 시간표에 다 못 들어간 항목</h4><ul>${overflowLines.map(l=>`<li>${escapeHtml(l)}</li>`).join('')}</ul></div>`
    : '';

  const mailLines = [];
  dayKeys.forEach(dk => {
    (weekData.mail[dk]||[]).forEach(m => mailLines.push(`${DAY_NAMES_FULL[dk]} · ${m.raw||''} · ${m.subject} → ${m.to}${m.summary ? ' · '+m.summary : ''}`));
  });
  el('mailArea').innerHTML = mailLines.length
    ? `<details class="mail-box"><summary>참고: 이번 주 관련 메일 ${mailLines.length}건</summary>${mailLines.map(l=>`<div class="mail-item">${escapeHtml(l)}</div>`).join('')}</details>`
    : '';
}

el('gridArea').addEventListener('input', e=>{
  const target = e.target;
  if(!target.classList.contains('cell-task')) return;
  const cellEl = target.closest('.cell-inner');
  const dk = cellEl.dataset.day, slotIdx = +cellEl.dataset.slotidx;
  const dayEntry = workingWeeks[currentWeekKey].days[dk];
  if(!dayEntry) return;
  dayEntry.slots[slotIdx].task = target.textContent.trim();
  dayEntry.slots[slotIdx].edited = true;
  dayEntry.slots[slotIdx].editedBy = currentRole;
  let badge = cellEl.querySelector('.cell-badge.edited');
  if(!badge){
    badge = document.createElement('span');
    badge.className = 'cell-badge edited';
    cellEl.appendChild(badge);
  }
  badge.className = 'cell-badge edited ' + currentRole;
  badge.textContent = `✎ ${currentRole==='mentor'?'선임자':'신입'} 수정`;
});

el('resetBtn').addEventListener('click', ()=>{
  if(!originalWeeks || !currentWeekKey) return;
  workingWeeks[currentWeekKey] = deepCopy(originalWeeks[currentWeekKey]);
  renderWeek(currentWeekKey);
});

el('genBtn').addEventListener('click', generate);

// AI 문서가 들어와 있으면 바로 일정표를 보여준다
if(AI_DOC_MODE){ generate(); }
</script>
</body>
</html>
'''


# =========================================================
# 0. Streamlit 기본 설정
# =========================================================
st.set_page_config(
    page_title="업무 인수인계 통합 지원 시스템",
    page_icon="📄",
    layout="wide",
)

st.markdown(
    """
    <style>
    :root {
        --navy: #17365D; --blue: #2F6FED; --blue-soft: #EAF1FF; --ink: #172033;
        --muted: #64748B; --line: #DCE5F0; --surface: #FFFFFF; --surface-soft: #F4F7FB;
    }
    html, body, [class*="css"] {
        font-family: "Pretendard", "Noto Sans KR", "Apple SD Gothic Neo", "Malgun Gothic", sans-serif;
    }
    .stApp { background: var(--surface-soft); color: var(--ink); }
    .block-container { max-width: 1440px; padding-top: 2rem; padding-bottom: 5rem; }
    #MainMenu, footer { visibility: hidden; }
    .handover-hero {
        background: radial-gradient(circle at 89% 16%, rgba(47,111,237,.17), transparent 27%),
                    linear-gradient(135deg, #FFFFFF 0%, #F8FBFF 100%);
        border: 1px solid #D8E3F2; border-radius: 24px; padding: 32px 36px; margin-bottom: 18px;
        box-shadow: 0 12px 34px rgba(23,54,93,.065);
    }
    .hero-eyebrow { color: var(--blue); font-size: .76rem; font-weight: 850; letter-spacing: .12em; margin-bottom: 9px; }
    .handover-hero h1 { color: var(--navy) !important; font-size: 2.25rem; line-height: 1.2; margin: 0 0 12px 0; letter-spacing: -.035em; font-weight: 850; }
    .handover-hero p { margin: 0; color: #526276; line-height: 1.75; font-size: 1rem; max-width: 980px; }
    .hero-chips { display: flex; gap: 8px; flex-wrap: wrap; margin-top: 20px; }
    .hero-chip { display: inline-flex; align-items: center; border: 1px solid #D7E3F5; background: #FFFFFF; color: #34506F; border-radius: 999px; padding: 7px 12px; font-size: .82rem; font-weight: 750; }
    .flow-strip { display: grid; grid-template-columns: repeat(4, minmax(0, 1fr)); gap: 10px; margin: 0 0 22px 0; }
    .flow-step { background: #FFFFFF; border: 1px solid var(--line); border-radius: 15px; padding: 14px 15px; color: #334A66; font-size: .88rem; font-weight: 750; text-align: center; box-shadow: 0 4px 14px rgba(23,54,93,.025); }
    div[data-baseweb="tab-list"] { gap: 3px; background: #FFFFFF; border: 1px solid var(--line); border-radius: 15px; padding: 5px; box-shadow: 0 4px 16px rgba(23,54,93,.03); }
    div[data-testid="stTabs"] button[role="tab"] { border-radius: 10px; padding-left: .82rem; padding-right: .82rem; font-weight: 750; color: #607086; }
    div[data-testid="stTabs"] button[role="tab"][aria-selected="true"] { background: var(--blue-soft); color: var(--navy); }
    h1, h2, h3 { color: var(--navy) !important; letter-spacing: -.025em; }
    div[data-testid="stMetric"] { background: #FFFFFF; border: 1px solid var(--line); border-radius: 16px; padding: 16px 17px; box-shadow: 0 5px 18px rgba(23,54,93,.03); }
    div[data-testid="stMetricLabel"] { color: #66778E; font-weight: 750; }
    div[data-testid="stMetricValue"] { color: var(--navy); font-weight: 850; }
    div[data-testid="stFileUploader"] { background: #FFFFFF; border: 1px dashed #B9C9DD; border-radius: 15px; padding: 7px; }
    div[data-testid="stTextInput"] input, div[data-testid="stTextArea"] textarea { border-radius: 10px !important; }
    div[data-testid="stButton"] > button, div[data-testid="stDownloadButton"] > button { border-radius: 10px; min-height: 42px; font-weight: 750; }
    div[data-testid="stAlert"] { border-radius: 13px; }
    div[data-testid="stExpander"] { background: #FFFFFF; border: 1px solid var(--line); border-radius: 13px; }
    div[data-testid="stDataFrame"] { border: 1px solid var(--line); border-radius: 13px; overflow: hidden; }
    div[data-testid="stChatMessage"] { background: #FFFFFF; border: 1px solid var(--line); border-radius: 16px; padding: 5px 10px; margin-bottom: 10px; }
    .section-note { background: #FFFFFF; border: 1px solid var(--line); border-left: 4px solid var(--blue); border-radius: 13px; padding: 14px 16px; margin: 6px 0 18px 0; color: #56667B; line-height: 1.65; font-size: .91rem; }
    @media (max-width: 900px) {
        .flow-strip { grid-template-columns: repeat(2, minmax(0, 1fr)); }
        .handover-hero { padding: 25px 23px; }
        .handover-hero h1 { font-size: 1.75rem; }
    }
    </style>
    """,
    unsafe_allow_html=True,
)

st.markdown(
    """
    <div class="handover-hero">
      <div class="hero-eyebrow">HANDOVER WORKSPACE · AI EDITION</div>
      <h1>업무 인수인계 통합 지원 시스템</h1>
      <p>
        형식이 제각각인 엑셀·회의록·메일·메모를 AI가 읽고 이해해, 후임자가 필요한 정보를 근거와 함께 조회하고
        첫 1~2주 온보딩 일정을 설계하며 표준 인수인계 문서를 생성할 수 있도록 지원합니다.
      </p>
      <div class="hero-chips">
        <span class="hero-chip">다양한 파일 AI 자동 인식</span>
        <span class="hero-chip">근거 기반 AI Q&amp;A</span>
        <span class="hero-chip">온보딩 일정 자동 설계</span>
        <span class="hero-chip">누락 항목 점검</span>
        <span class="hero-chip">DOCX 인수인계서 생성</span>
        <span class="hero-chip">오류 자동 감지·신고</span>
      </div>
    </div>
    <div class="flow-strip">
      <div class="flow-step">① 업무자료 통합</div>
      <div class="flow-step">② 업무 조회·Q&amp;A</div>
      <div class="flow-step">③ 온보딩 일정 설계</div>
      <div class="flow-step">④ 인수인계서 생성</div>
    </div>
    """,
    unsafe_allow_html=True,
)

# AI 연결 상태 표시줄 (내용은 화면 맨 아래에서 채움 → 이번 실행의 사용량까지 반영)
ai_status_slot = st.empty()


# =========================================================
# 1. 자동 오류 감지 / 오류 리포트 / 사용자 행동 로그
# =========================================================
ERROR_REPORT_PATH = Path("data/error_reports.jsonl")
SLOW_RESPONSE_SECONDS = 30.0  # AI 응답은 몇 초~십수 초 걸리므로 기준을 30초로

SCREEN_FUNCTION_MAP = {
    "🏠 대시보드": "후임자 업무 대시보드",
    "📂 Excel 업무자료": "업무자료 업로드 및 자동 분류",
    "💬 후임자 Q&A": "후임자 질문 검색",
    "📅 온보딩 일과표": "온보딩 일과표 생성 및 수정",
    "🤖 업무메모 자동 인수인계": "업무메모 자동 인수인계",
    "✍️ 직접 작성": "인수인계서 직접 작성",
    "🛠 관리자 오류함": "오류 신고 관리",
}
MAIN_SCREEN_OPTIONS = list(SCREEN_FUNCTION_MAP.keys())


def _mask_sensitive_text(value):
    """오류 신고 전에 주민등록번호처럼 명확한 민감 패턴을 간단히 마스킹."""
    if value is None:
        return ""
    text = str(value)
    text = re.sub(r"\b(\d{6})[- ]?([1-4])\d{6}\b", r"\1-*******", text)
    # API 키가 오류 메시지에 섞여 들어가지 않도록 마스킹
    text = re.sub(r"sk-[A-Za-z0-9_\-]{8,}", "sk-****", text)
    return text


def _load_error_reports():
    if not ERROR_REPORT_PATH.exists():
        return []
    reports = []
    try:
        with ERROR_REPORT_PATH.open("r", encoding="utf-8") as f:
            for line in f:
                line = line.strip()
                if not line:
                    continue
                try:
                    reports.append(json.loads(line))
                except json.JSONDecodeError:
                    continue
    except OSError:
        return []
    return reports[-100:]


def init_error_reporting():
    defaults = {
        "error_action_logs": [],
        "last_screen": "🏠 대시보드",
        "last_function": "앱 이용",
        "last_input": "",
        "last_response_time": None,
        "last_input_by_screen": {},
        "last_response_time_by_screen": {},
        "last_error": None,
        "last_uploaded_files": [],
        "pending_auto_error_dialog": False,
    }
    for key, value in defaults.items():
        if key not in st.session_state:
            st.session_state[key] = value
    if "error_reports" not in st.session_state:
        st.session_state["error_reports"] = _load_error_reports()


def track_action(screen, function, action, input_value="", metadata=None):
    """최근 사용자 행동을 최대 20개까지 세션에 기록."""
    event = {
        "time": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "screen": screen,
        "function": function,
        "action": action,
        "input": _mask_sensitive_text(input_value),
        "metadata": metadata or {},
    }
    logs = list(st.session_state.get("error_action_logs", []))
    logs.append(event)
    st.session_state["error_action_logs"] = logs[-20:]
    st.session_state["last_screen"] = screen
    st.session_state["last_function"] = function

    if input_value:
        masked_input = _mask_sensitive_text(input_value)
        st.session_state["last_input"] = masked_input
        input_by_screen = dict(st.session_state.get("last_input_by_screen", {}))
        input_by_screen[screen] = masked_input
        st.session_state["last_input_by_screen"] = input_by_screen


def _queue_auto_error_dialog():
    st.session_state["pending_auto_error_dialog"] = True


def _store_error(screen, function, error_info, input_value, response_time, action_label):
    st.session_state["last_error"] = error_info
    st.session_state["last_response_time"] = response_time
    response_by_screen = dict(st.session_state.get("last_response_time_by_screen", {}))
    response_by_screen[screen] = response_time
    st.session_state["last_response_time_by_screen"] = response_by_screen
    track_action(
        screen=screen,
        function=function,
        action=action_label,
        input_value=input_value,
        metadata={"error_type": error_info["type"], "error_message": error_info["message"]},
    )
    _queue_auto_error_dialog()


def record_error(screen, function, error, code="APP_ERROR", input_value="", response_time=None):
    """Python 예외를 자동 감지하고 오류 리포트 초안을 준비."""
    error_info = {
        "code": code,
        "type": type(error).__name__,
        "message": _mask_sensitive_text(str(error)),
        "occurred_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "response_time": response_time,
        "traceback": _mask_sensitive_text(traceback.format_exc()),
    }
    _store_error(screen, function, error_info, input_value, response_time, f"오류 자동 감지 ({code})")


def record_system_issue(screen, function, code, message, issue_type="SYSTEM_WARNING",
                        input_value="", response_time=None):
    """예외가 없어도 응답 지연처럼 시스템이 판단 가능한 이상 상태를 자동 감지."""
    error_info = {
        "code": code,
        "type": issue_type,
        "message": _mask_sensitive_text(message),
        "occurred_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "response_time": response_time,
        "traceback": "",
    }
    _store_error(screen, function, error_info, input_value, response_time, f"이상 상태 자동 감지 ({code})")


def _action_lines_for(screen):
    logs = [e for e in st.session_state.get("error_action_logs", []) if e.get("screen") == screen][-8:]
    lines = []
    for idx, event in enumerate(logs, start=1):
        line = f"{idx}. [{event.get('time', '')}] {event.get('screen', '')} · {event.get('action', '')}"
        if event.get("input"):
            line += f" · 입력: {event.get('input', '')}"
        lines.append(line)
    return lines


def build_error_report_draft():
    """자동 오류 신고 초안. 오류가 발생한 화면의 기록만 사용합니다."""
    current_screen = st.session_state.get("last_screen", "화면 미확인")
    current_function = st.session_state.get("last_function", "기능 미확인")
    last_error = st.session_state.get("last_error") or {}
    last_input = st.session_state.get("last_input_by_screen", {}).get(current_screen, "")
    response_time = last_error.get(
        "response_time",
        st.session_state.get("last_response_time_by_screen", {}).get(current_screen),
    )
    action_lines = _action_lines_for(current_screen)
    error_code = last_error.get("code", "AUTO_DETECT")
    error_message = last_error.get("message", "자동 감지된 오류 메시지 없음")

    summary_parts = [f"{current_screen}에서 ", f"{current_function} 기능 이용 중 시스템이 오류를 자동 감지했습니다."]
    if last_input:
        summary_parts.append(f" 최근 입력/검색어는 '{last_input}'입니다.")
    if last_error:
        summary_parts.append(f" 감지된 오류는 {error_code}이며, 메시지는 '{error_message}'입니다.")
    if response_time is not None:
        summary_parts.append(f" 마지막 측정 응답 시간은 {float(response_time):.2f}초입니다.")

    return {
        "screen": current_screen,
        "function": current_function,
        "last_input": last_input,
        "response_time": response_time,
        "error_code": error_code,
        "error_type": last_error.get("type", ""),
        "error_message": error_message,
        "action_history": "\n".join(action_lines) if action_lines else "현재 화면에서 기록된 최근 동작이 없습니다.",
        "auto_summary": "".join(summary_parts),
        "technical_traceback": last_error.get("traceback", ""),
    }


def save_error_report(report):
    ERROR_REPORT_PATH.parent.mkdir(parents=True, exist_ok=True)
    with ERROR_REPORT_PATH.open("a", encoding="utf-8") as f:
        f.write(json.dumps(report, ensure_ascii=False) + "\n")
    reports = list(st.session_state.get("error_reports", []))
    reports.append(report)
    st.session_state["error_reports"] = reports[-100:]


def delete_error_report(report_id):
    """선택한 오류 신고 1건을 세션과 JSONL 저장 파일에서 함께 삭제."""
    if not report_id:
        return False
    reports = list(st.session_state.get("error_reports", []))
    remaining = [r for r in reports if r.get("report_id") != report_id]
    if len(remaining) == len(reports):
        return False
    st.session_state["error_reports"] = remaining[-100:]
    ERROR_REPORT_PATH.parent.mkdir(parents=True, exist_ok=True)
    temp_path = ERROR_REPORT_PATH.with_suffix(".jsonl.tmp")
    with temp_path.open("w", encoding="utf-8") as f:
        for report in remaining[-100:]:
            f.write(json.dumps(report, ensure_ascii=False) + "\n")
    temp_path.replace(ERROR_REPORT_PATH)
    return True


@st.dialog("⚠️ 시스템 오류 자동 감지", width="large")
def show_error_report_dialog():
    draft = build_error_report_draft()
    st.warning("시스템이 오류 또는 이상 상태를 자동으로 감지했습니다.")
    st.caption(
        "오류 발생 당시의 화면, 기능, 입력값, 최근 동작, 에러 코드와 응답 시간을 자동으로 수집했습니다. "
        "내용을 확인하고 필요한 경우 추가 의견만 작성한 뒤 관리자에게 전송하세요."
    )
    c1, c2 = st.columns(2)
    with c1:
        st.text_input("현재 화면", value=draft["screen"], disabled=True)
    with c2:
        st.text_input("사용 기능", value=draft["function"], disabled=True)
    st.text_input("최근 입력/검색어", value=draft["last_input"] or "기록 없음", disabled=True)
    c3, c4 = st.columns(2)
    with c3:
        st.text_input("에러 코드", value=draft["error_code"], disabled=True)
    with c4:
        response_label = (
            f"{float(draft['response_time']):.2f}초" if draft["response_time"] is not None else "측정 기록 없음"
        )
        st.text_input("최근 응답 시간", value=response_label, disabled=True)
    st.text_area("최근 동작 순서", value=draft["action_history"], height=160, disabled=True)
    auto_summary = st.text_area(
        "자동 작성된 오류 내용", value=draft["auto_summary"], height=130,
        help="필요하면 사용자가 직접 수정할 수 있습니다.", key="error_report_auto_summary",
    )
    additional_note = st.text_area(
        "추가로 전달할 내용 (선택)",
        placeholder="예: 같은 버튼을 두 번 눌러도 동일한 오류가 발생했습니다.",
        key="error_report_additional_note",
    )
    st.caption("🔒 주민등록번호 형태의 값과 API 키는 자동 마스킹합니다.")

    if st.button("📨 관리자에게 전송", type="primary", use_container_width=True, key="submit_error_report"):
        report = {
            "report_id": "ERR-" + uuid.uuid4().hex[:8].upper(),
            "reported_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            "screen": draft["screen"],
            "function": draft["function"],
            "last_input": draft["last_input"],
            "response_time": draft["response_time"],
            "error_code": draft["error_code"],
            "error_type": draft["error_type"],
            "error_message": draft["error_message"],
            "action_history": draft["action_history"],
            "auto_summary": _mask_sensitive_text(auto_summary),
            "additional_note": _mask_sensitive_text(additional_note),
            "screenshot": "",
            "status": "미처리",
            "technical_traceback": draft["technical_traceback"],
        }
        try:
            save_error_report(report)
            st.session_state["pending_auto_error_dialog"] = False
            st.session_state["last_error"] = None
            st.success(f"✅ 오류 리포트가 관리자 오류함에 접수되었습니다. 신고번호: {report['report_id']}")
        except Exception as e:
            st.error(f"신고 저장 중 오류가 발생했습니다: {e}")


@st.dialog("⚠️ 오류 신고", width="large")
def show_manual_error_report_dialog():
    """수동 오류 신고 팝업. 현재 선택 화면의 기록만 표시합니다."""
    tracked_tab = st.session_state.get("main_tabs")
    fallback_screen = st.session_state.get("last_screen", "🏠 대시보드")
    if tracked_tab in MAIN_SCREEN_OPTIONS:
        default_screen = tracked_tab
    elif fallback_screen in MAIN_SCREEN_OPTIONS:
        default_screen = fallback_screen
    else:
        default_screen = "🏠 대시보드"

    selected_screen = st.session_state.get("manual_report_screen_select", default_screen)
    if selected_screen not in MAIN_SCREEN_OPTIONS:
        selected_screen = default_screen
    default_index = MAIN_SCREEN_OPTIONS.index(selected_screen)

    st.info(
        "시스템이 자동으로 감지하지 못한 이상 현상이 있다면 직접 신고할 수 있습니다. "
        "현재 화면의 기록만 자동으로 첨부됩니다."
    )
    current_screen = st.selectbox(
        "현재 화면", MAIN_SCREEN_OPTIONS, index=default_index,
        key="manual_report_screen_select", help="현재 보고 있는 화면이 맞는지 확인해주세요.",
    )
    current_function = SCREEN_FUNCTION_MAP[current_screen]
    st.markdown("**최근 사용 기능**")
    st.code(current_function, language=None)

    last_input = st.session_state.get("last_input_by_screen", {}).get(current_screen, "")
    last_response_time = st.session_state.get("last_response_time_by_screen", {}).get(current_screen)
    st.markdown("**최근 입력/검색어**")
    st.code(last_input or "현재 화면에서 기록된 입력 없음", language=None)
    if last_response_time is not None:
        st.caption(f"최근 측정 응답 시간: {float(last_response_time):.2f}초")

    action_lines = _action_lines_for(current_screen)
    st.markdown("**최근 동작 순서**")
    st.code("\n".join(action_lines) if action_lines else "현재 화면에서 기록된 최근 동작이 없습니다.", language=None)

    user_note = st.text_area(
        "어떤 문제가 있었나요?",
        placeholder="예: 일정표 생성 버튼을 눌렀는데 결과가 표시되지 않았어요.",
        height=120, key="manual_error_report_note",
    )
    st.caption("🔒 주민등록번호 형태의 값과 API 키는 신고 저장 전에 자동 마스킹됩니다.")

    if st.button("📨 관리자에게 신고", type="primary", use_container_width=True, key="manual_error_report_submit"):
        if not user_note.strip():
            st.warning("신고할 내용을 간단히 입력해주세요.")
            return
        report = {
            "report_id": "ERR-" + uuid.uuid4().hex[:8].upper(),
            "reported_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            "screen": current_screen,
            "function": current_function,
            "last_input": _mask_sensitive_text(last_input),
            "response_time": last_response_time,
            "error_code": "USER_REPORT",
            "error_type": "USER_REPORTED_ISSUE",
            "error_message": _mask_sensitive_text(user_note),
            "action_history": "\n".join(action_lines) if action_lines else "현재 화면에서 기록된 최근 동작이 없습니다.",
            "auto_summary": f"{current_screen}에서 {current_function} 기능 사용 중 사용자가 이상 현상을 직접 신고했습니다.",
            "additional_note": _mask_sensitive_text(user_note),
            "screenshot": "",
            "status": "미처리",
            "technical_traceback": "",
        }
        try:
            save_error_report(report)
            st.session_state["last_screen"] = current_screen
            st.session_state["last_function"] = current_function
            st.success(f"✅ 신고가 접수되었습니다. 신고번호: {report['report_id']}")
        except Exception as e:
            st.error(f"신고 저장 중 오류가 발생했습니다: {e}")


def render_floating_error_button():
    """메인 화면 전체에 단 하나의 플로팅 오류 신고 버튼만 표시합니다."""
    st.markdown(
        """
        <style>
        .st-key-floating_error_report_global { position: fixed; right: 24px; bottom: 24px; z-index: 999999; width: auto !important; }
        .st-key-floating_error_report_global button { border-radius: 999px !important; padding: 0.65rem 1rem !important; box-shadow: 0 4px 18px rgba(0, 0, 0, 0.20); font-weight: 700; }
        </style>
        """,
        unsafe_allow_html=True,
    )
    with st.container(key="floating_error_report_global"):
        if st.button("⚠️ 오류 신고", key="floating_error_report_button_global",
                     help="현재 화면의 오류 또는 이상 현상을 신고합니다."):
            active_screen = st.session_state.get("main_tabs", st.session_state.get("last_screen", "🏠 대시보드"))
            if active_screen not in SCREEN_FUNCTION_MAP:
                active_screen = "🏠 대시보드"
            st.session_state["last_screen"] = active_screen
            st.session_state["last_function"] = SCREEN_FUNCTION_MAP[active_screen]
            st.session_state["manual_report_screen_select"] = active_screen
            show_manual_error_report_dialog()


init_error_reporting()


# =========================================================
# 1-1. AI 공통 도우미
# =========================================================
def content_hash(*parts):
    h = hashlib.sha256()
    for part in parts:
        if isinstance(part, bytes):
            h.update(part)
        else:
            h.update(str(part).encode("utf-8"))
        h.update(b"\x00")
    return h.hexdigest()[:16]


def ai_try(tag, key_hash, fn, screen, function, code, input_value=""):
    """
    AI 기능을 한 번 시도한다. 실패하면 오류를 기록하고 None 을 돌려준다.
    같은 자료로 이미 실패했다면 재시도하지 않는다(화면이 다시 그려질 때마다 오류창이 뜨는 것 방지).
    """
    flag = f"_ai_failed_{tag}_{key_hash}"
    if not ai.ai_available() or st.session_state.get(flag):
        return None
    try:
        return fn()
    except Exception as e:
        st.session_state[flag] = True
        record_error(screen=screen, function=function, error=e, code=code, input_value=input_value)
        return None


def get_shared_docs():
    """온보딩 탭에 올린 회의록·캘린더·메일 문서를 [(파일명, 텍스트)] 로 반환 (Q&A·인수인계서에서도 사용)."""
    docs = []
    for f in st.session_state.get("onboarding_docs") or []:
        try:
            docs.append((f.name, ai.extract_text(f.name, f.getvalue())))
        except Exception:
            continue
    return docs


def get_excel_sources():
    """엑셀 탭에 올린 파일을 [(파일명, 텍스트)] 로 반환 (인수인계서 종합 생성용)."""
    sources = []
    for f in st.session_state.get("excel_upload") or []:
        try:
            sources.append((f.name, ai.extract_text(f.name, f.getvalue())))
        except Exception:
            continue
    return sources


def render_ai_status():
    ok, reason = ai.ai_status()
    if ok:
        usage = ai.usage_summary()
        ai_status_slot.success(
            f"🤖 **AI 연결됨** · 모델 `{ai.get_model()}` · 이번 접속 중 AI 호출 {usage['calls']}회 "
            f"· 예상 비용 약 {usage['krw']:,}원 (같은 자료·같은 질문은 다시 요금이 나가지 않아요)"
        )
    else:
        ai_status_slot.warning(f"⚙️ **AI 미연결 — 기존 규칙 기반 방식으로 동작합니다.** {reason}")


# =========================================================
# 2. 공통 유틸리티
# =========================================================
def default_rows(columns, n=3):
    return pd.DataFrame([{c: "" for c in columns} for _ in range(n)])


def clean_records(df):
    """빈 행을 제거하고 dict 리스트로 변환."""
    if df is None:
        return []
    df = df.fillna("")
    records = []
    for row in df.to_dict(orient="records"):
        normalized = {str(k): str(v).strip() for k, v in row.items()}
        if any(normalized.values()):
            records.append(normalized)
    return records


def safe_value(value):
    """NaN/None 등을 빈 문자열로 정리."""
    if value is None:
        return ""
    try:
        if pd.isna(value):
            return ""
    except Exception:
        pass
    return str(value).strip()


def read_excel_smart(uploaded_file):
    """(AI 미사용 시 예비) 제목 행이 위에 있어도 실제 헤더 행을 찾아 읽습니다."""
    raw = pd.read_excel(uploaded_file, header=None)
    known_headers = {"일자", "업무", "프로젝트/현장", "현재 단계", "구분", "회사/부서", "시스템/자산", "유형"}
    header_index = 0
    for i in range(min(len(raw), 10)):
        values = {str(v).strip() for v in raw.iloc[i].tolist() if pd.notna(v)}
        if len(known_headers & values) >= 2:
            header_index = i
            break
    uploaded_file.seek(0)
    df = pd.read_excel(uploaded_file, header=header_index)
    df = df.dropna(how="all").dropna(axis=1, how="all")
    df.columns = [str(c).strip() for c in df.columns]
    df = df.reset_index(drop=True)
    df.attrs["excel_header_row"] = header_index + 1
    return df


def classify_excel_files(excel_data):
    """(AI 미사용 시 예비) 파일 이름으로 분류."""
    schedule_df = project_df = contact_df = asset_df = None
    unknown_files = []
    for file_name, df in excel_data.items():
        normalized_name = file_name.replace(" ", "")
        if "업무일정" in normalized_name:
            schedule_df = df
        elif "프로젝트" in normalized_name:
            project_df = df
        elif "담당자" in normalized_name:
            contact_df = df
        elif "계정" in normalized_name or "권한" in normalized_name or "자산" in normalized_name:
            asset_df = df
        else:
            unknown_files.append(file_name)
    return schedule_df, project_df, contact_df, asset_df, unknown_files


def rule_normalize_excel(uploaded_files):
    """AI 없이 기존 규칙(파일 이름 + 고정 헤더)으로 읽기."""
    excel_data = {}
    for uploaded_file in uploaded_files:
        try:
            if uploaded_file.name.lower().endswith(".csv"):
                excel_data[uploaded_file.name] = pd.read_csv(uploaded_file)
            else:
                excel_data[uploaded_file.name] = read_excel_smart(uploaded_file)
            excel_data[uploaded_file.name].attrs["source_file"] = uploaded_file.name
        except Exception as e:
            record_error(screen="📂 Excel 업무자료", function="Excel 파일 읽기", error=e,
                         code="EXCEL_READ_ERROR", input_value=uploaded_file.name)
            st.error(f"❌ {uploaded_file.name} 읽기 실패 · 시스템이 오류를 자동 감지했습니다.")
    s, p, c, a, unknown = classify_excel_files(excel_data)
    result = {
        "schedule": ai.ensure_canonical_columns(s, "schedule"),
        "project": ai.ensure_canonical_columns(p, "project"),
        "contact": ai.ensure_canonical_columns(c, "contact"),
        "asset": ai.ensure_canonical_columns(a, "asset"),
        "unknown": unknown,
        "mapping": [],
    }
    return result


# =========================================================
# 3. DOCX 생성 함수
# =========================================================
def set_cell_text(cell, text, bold=False, size=9):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(str(text) if text is not None else "")
    run.bold = bold
    run.font.name = "맑은 고딕"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    run.font.size = Pt(size)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def style_document(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "맑은 고딕"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    normal.font.size = Pt(9)


def add_title(doc, title):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    run.bold = True
    run.font.name = "맑은 고딕"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    run.font.size = Pt(18)


def add_section_heading(doc, title):
    p = doc.add_paragraph()
    run = p.add_run(title)
    run.bold = True
    run.font.name = "맑은 고딕"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    run.font.size = Pt(11)


def add_key_value_table(doc, pairs, cols=2):
    rows = (len(pairs) + cols - 1) // cols
    table = doc.add_table(rows=rows, cols=cols * 2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    idx = 0
    for r in range(rows):
        for c in range(cols):
            if idx < len(pairs):
                key, value = pairs[idx]
                set_cell_text(table.cell(r, c * 2), key, bold=True)
                set_cell_text(table.cell(r, c * 2 + 1), value)
                idx += 1
            else:
                set_cell_text(table.cell(r, c * 2), "")
                set_cell_text(table.cell(r, c * 2 + 1), "")
    return table


def add_records_table(doc, records, columns):
    table = doc.add_table(rows=1, cols=len(columns))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, col in enumerate(columns):
        set_cell_text(table.rows[0].cells[i], col, bold=True)
    if records:
        for record in records:
            cells = table.add_row().cells
            for i, col in enumerate(columns):
                set_cell_text(cells[i], record.get(col, ""))
    else:
        cells = table.add_row().cells
        for i in range(len(columns)):
            set_cell_text(cells[i], "")
    return table


def add_detail_table(doc, details):
    table = doc.add_table(rows=len(details), cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (key, value) in enumerate(details):
        set_cell_text(table.cell(i, 0), key, bold=True)
        set_cell_text(table.cell(i, 1), value)
    return table


DETAIL_DOC_FIELDS = [
    "업무 개요", "목적 / 성과지표", "진행 중 프로젝트 현황", "정기 업무", "비정기 업무",
    "주요 일정 / 마감", "관련 시스템 / 계정 / 권한", "관련 담당자 / 연락처", "협업 부서",
    "특이사항 / 주의사항", "리스크 / 미해결 이슈", "참고 파일 경로 / 문서 링크",
    "후임자 숙지 필요사항", "인수인계 완료 여부",
]


def build_docx(data):
    doc = Document()
    style_document(doc)
    add_title(doc, "업무 인수인계서")
    add_key_value_table(doc, [
        ("기관명", data["meta"]["기관명"]), ("부서명", data["meta"]["부서명"]),
        ("문서번호", data["meta"]["문서번호"]), ("보존기간", data["meta"]["보존기간"]),
    ])

    add_section_heading(doc, "I. 기본 정보")
    add_key_value_table(doc, [
        ("소속 부서", data["basic"]["소속 부서"]), ("직위 / 직책", data["basic"]["직위 / 직책"]),
        ("인계자 성명", data["basic"]["인계자 성명"]), ("인수자 성명", data["basic"]["인수자 성명"]),
        ("작성일", data["basic"]["작성일"]), ("인수인계 완료 예정일", data["basic"]["인수인계 완료 예정일"]),
    ])

    add_section_heading(doc, "II. 담당업무 개요")
    add_records_table(doc, data["tasks"], ["주요 업무", "업무 목적", "업무 프로세스", "우선순위", "비고"])

    add_section_heading(doc, "III. 업무 상세")
    for i, item in enumerate(data["details"], start=1):
        p = doc.add_paragraph()
        run = p.add_run(f"{i}) {item.get('업무명', '업무 상세')}")
        run.bold = True
        run.font.name = "맑은 고딕"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
        run.font.size = Pt(10)
        add_detail_table(doc, [(field, item.get(field, "")) for field in DETAIL_DOC_FIELDS])

    add_section_heading(doc, "IV. 주요 일정 및 대외 커뮤니케이션")
    p = doc.add_paragraph()
    p.add_run("1. 긴급 업무 일정").bold = True
    add_records_table(doc, data["urgent_schedule"], ["일자", "내용", "대응 방법", "담당"])
    p = doc.add_paragraph()
    p.add_run("2. 향후 1개월 주요 일정").bold = True
    add_records_table(doc, data["monthly_schedule"], ["일자", "일정 내용", "비고"])
    add_detail_table(doc, [("3. 대외 커뮤니케이션 유의사항", data["communication_note"])])

    add_section_heading(doc, "V. 계정 · 권한 · 자산 인계")
    add_records_table(doc, data["assets"], ["시스템 / 자산명", "유형", "권한 수준", "인계 방법", "상태", "비고"])

    add_section_heading(doc, "VI. 인수인계 체크리스트")
    add_records_table(doc, data["checklist"], ["체크", "항목", "확인일", "확인자", "비고"])

    add_section_heading(doc, "VII. 서명")
    add_key_value_table(doc, [
        ("인계자 성명", data["signatures"]["인계자 성명"]), ("인계자 서명일", data["signatures"]["인계자 서명일"]),
        ("인수자 성명", data["signatures"]["인수자 성명"]), ("인수자 서명일", data["signatures"]["인수자 서명일"]),
        ("확인자(팀장 등) 성명", data["signatures"]["확인자 성명"]), ("확인자 서명일", data["signatures"]["확인자 서명일"]),
    ])

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("본 문서는 업무 연속성을 위해 작성된 내부 인수인계 자료입니다. 외부 반출 시 소속 부서의 사전 승인이 필요합니다.")
    run.font.name = "맑은 고딕"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    run.font.size = Pt(8)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


# =========================================================
# 4. 인수인계 준비도 계산
# =========================================================
def calculate_readiness(reviewed_data):
    missing_items = []
    check_items = []

    def check_value(label, value):
        if value is not None and str(value).strip() != "":
            check_items.append(True)
        else:
            check_items.append(False)
            missing_items.append(label)

    for key in ["소속 부서", "직위 / 직책", "인계자 성명", "인수자 성명", "작성일", "인수인계 완료 예정일"]:
        check_value(key, reviewed_data["basic"].get(key, ""))

    if reviewed_data["tasks"]:
        check_items.append(True)
    else:
        check_items.append(False)
        missing_items.append("담당업무 개요")

    if not reviewed_data["details"]:
        missing_items.append("업무 상세")
        check_items.extend([False] * 7)
    else:
        for i, detail in enumerate(reviewed_data["details"], start=1):
            check_value(f"업무 {i} - 업무명", detail.get("업무명", ""))
            check_value(f"업무 {i} - 업무 개요", detail.get("업무 개요", ""))
            check_value(f"업무 {i} - 주요 일정 / 마감", detail.get("주요 일정 / 마감", ""))
            check_value(f"업무 {i} - 관련 담당자 / 연락처", detail.get("관련 담당자 / 연락처", ""))
            check_value(f"업무 {i} - 리스크 / 미해결 이슈", detail.get("리스크 / 미해결 이슈", ""))
            check_value(f"업무 {i} - 참고 파일", detail.get("참고 파일 경로 / 문서 링크", ""))
            check_value(f"업무 {i} - 후임자 숙지사항", detail.get("후임자 숙지 필요사항", ""))

    for key, label in [("urgent_schedule", "긴급 업무 일정"), ("monthly_schedule", "향후 1개월 주요 일정"),
                       ("assets", "계정 · 권한 · 자산 인계")]:
        if reviewed_data[key]:
            check_items.append(True)
        else:
            check_items.append(False)
            missing_items.append(label)

    total_count = len(check_items)
    completed_count = sum(check_items)
    readiness = int(completed_count / total_count * 100) if total_count else 0
    return readiness, completed_count, missing_items


def show_readiness(reviewed_data):
    readiness, completed_count, missing_items = calculate_readiness(reviewed_data)
    st.subheader("📊 인수인계 준비도")
    col1, col2, col3 = st.columns(3)
    col1.metric("인수인계 준비도", f"{readiness}%")
    col2.metric("작성 완료 항목", f"{completed_count}개")
    col3.metric("보완 필요 항목", f"{len(missing_items)}개")
    st.progress(readiness / 100)
    if readiness >= 90:
        st.success("✅ 인수인계 준비 상태가 매우 좋습니다.")
    elif readiness >= 70:
        st.warning("⚠️ 일부 항목의 보완이 필요합니다.")
    else:
        st.error("🚨 인수인계에 필요한 정보가 부족합니다.")
    if missing_items:
        with st.expander("⚠️ 보완이 필요한 항목 보기", expanded=True):
            for item in missing_items:
                st.write(f"• {item}")
    else:
        st.success("🎉 필수 인수인계 항목이 모두 작성되었습니다!")


# =========================================================
# 5. Excel 업무자료 -> 온보딩 일과표 HTML 연동
# =========================================================
def _first_existing_value(row, candidates, default=""):
    for column in candidates:
        if column in row.index:
            value = safe_value(row.get(column, ""))
            if value:
                return value
    return default


def _date_from(row, candidates):
    raw = _first_existing_value(row, candidates)
    return _format_dashboard_date(raw) if raw else ""


def _df_to_onboarding_data(schedule_df, project_df, contact_df, asset_df):
    """업로드한 업무자료를 일정표 HTML의 JavaScript 배열 구조로 변환."""
    data = {k: None for k in ["schedule_tasks", "project_deadline_tasks", "asset_deadline_tasks",
                              "project_overview", "asset_overview", "contacts"]}

    if schedule_df is not None:
        tasks = []
        for _, row in schedule_df.iterrows():
            title = _first_existing_value(row, ["업무", "업무명", "일정", "내용"])
            if not title:
                continue
            tasks.append({
                "date": _date_from(row, ["일자", "날짜", "마감", "마감일"]),
                "title": title,
                "priority": _first_existing_value(row, ["우선순위", "중요도"], "중"),
                "nextAction": _first_existing_value(row, ["다음 조치", "다음조치", "다음 액션", "다음액션"]),
                "note": _first_existing_value(row, ["비고", "주의사항", "메모"]),
                "source": "업무일정 · " + schedule_df.attrs.get("source_file", "Excel"),
            })
        data["schedule_tasks"] = tasks

    if project_df is not None:
        deadlines, overview = [], []
        source = "프로젝트 현황 · " + project_df.attrs.get("source_file", "Excel")
        for _, row in project_df.iterrows():
            name = _first_existing_value(row, ["프로젝트/현장", "프로젝트", "현장", "프로젝트명"])
            if not name:
                continue
            deadline = _date_from(row, ["마감일", "마감", "완료 예정일", "완료예정일"])
            stage = _first_existing_value(row, ["현재 단계", "단계", "상태"])
            progress = _first_existing_value(row, ["진행률", "진척률"])
            next_action = _first_existing_value(row, ["다음 액션", "다음액션", "다음 조치", "다음조치"])
            risk = _first_existing_value(row, ["리스크/이슈", "리스크", "이슈", "특이사항"])
            priority = _first_existing_value(row, ["우선순위", "중요도"], "상")
            if deadline:
                deadlines.append({"date": deadline, "title": f"[마감] {name}", "priority": priority,
                                  "nextAction": next_action, "note": risk, "source": source})
            parts = []
            if stage:
                parts.append(stage)
            if progress:
                parts.append(progress if "%" in str(progress) else f"진행률 {progress}%")
            if next_action:
                parts.append(f"다음액션: {next_action}")
            if risk:
                parts.append(f"이슈: {risk}")
            overview.append({"title": f"[개요] {name} 현황 파악", "detail": " · ".join(parts),
                             "badge": "개요", "source": source})
        data["project_deadline_tasks"] = deadlines
        data["project_overview"] = overview

    if contact_df is not None:
        contacts = []
        source = "담당자 연락망 · " + contact_df.attrs.get("source_file", "Excel")
        for _, row in contact_df.iterrows():
            name = _first_existing_value(row, ["성명", "이름", "담당자", "담당자명"])
            if not name:
                continue
            position = _first_existing_value(row, ["직책", "직급"])
            organization = _first_existing_value(row, ["회사/부서", "회사", "부서", "소속"])
            related = _first_existing_value(row, ["관련 업무", "관련업무", "담당 업무", "담당업무"])
            caution = _first_existing_value(row, ["커뮤니케이션 유의사항", "유의사항", "주의사항", "비고"])
            sub = ", ".join(v for v in [position, organization] if v)
            label = f"{name}({sub})" if sub else name
            parts = []
            if related:
                parts.append(f"관련 업무: {related}")
            if caution:
                parts.append(f"유의사항: {caution}")
            contacts.append({"title": f"{label} 컨택포인트 파악", "detail": " · ".join(parts),
                             "badge": "소개", "source": source})
        data["contacts"] = contacts

    if asset_df is not None:
        deadlines, overview = [], []
        source = "계정·권한·자산 · " + asset_df.attrs.get("source_file", "Excel")
        for _, row in asset_df.iterrows():
            name = _first_existing_value(row, ["시스템/자산", "시스템", "자산", "시스템/자산명"])
            if not name:
                continue
            deadline = _date_from(row, ["완료 목표일", "완료목표일", "마감일", "마감"])
            status = _first_existing_value(row, ["상태", "인계 상태", "인계상태"])
            method = _first_existing_value(row, ["인계 방법", "인계방법"])
            caution = _first_existing_value(row, ["주의사항", "유의사항", "비고"])
            priority = _first_existing_value(row, ["우선순위", "중요도"], "상")
            # 이미 완료된 인계는 마감 일정에 넣지 않음
            if deadline and status != "완료":
                deadlines.append({"date": deadline, "title": f"[마감] {name} 인계", "priority": priority,
                                  "nextAction": method, "note": caution, "source": source})
            parts = []
            if status:
                parts.append(f"상태: {status}")
            if method:
                parts.append(f"인계방법: {method}")
            if caution:
                parts.append(f"주의사항: {caution}")
            overview.append({"title": f"[개요] {name} 인계 상태 확인", "detail": " · ".join(parts),
                             "badge": "개요", "source": source})
        data["asset_deadline_tasks"] = deadlines
        data["asset_overview"] = overview

    return data


def _js_json(value):
    # </script> 로 HTML이 끊기지 않도록 이스케이프
    return json.dumps(value, ensure_ascii=False).replace("</", "<\\/")


def _replace_js_array(html, constant_name, value):
    """HTML 안의 const NAME = [ ... ]; 배열을 교체. value가 None이면 데모 데이터 유지."""
    if value is None:
        return html
    replacement = f"const {constant_name} = {_js_json(value)};"
    pattern = rf"const\s+{re.escape(constant_name)}\s*=\s*\[.*?\];"
    updated, count = re.subn(pattern, lambda m: replacement, html, count=1, flags=re.DOTALL)
    if count != 1:
        raise ValueError(f"일정표 HTML 데이터 영역을 찾지 못했습니다: {constant_name}")
    return updated


def _replace_js_const(html, constant_name, value):
    """한 줄짜리 const NAME = ...; 값을 교체."""
    replacement = f"const {constant_name} = {_js_json(value)};"
    pattern = rf"^const\s+{re.escape(constant_name)}\s*=.*;$"
    updated, count = re.subn(pattern, lambda m: replacement, html, count=1, flags=re.MULTILINE)
    if count != 1:
        raise ValueError(f"일정표 HTML 설정 영역을 찾지 못했습니다: {constant_name}")
    return updated


def build_onboarding_schedule_html(schedule_df, project_df, contact_df, asset_df, ai_items=None):
    """원본 HTML/JS 로직은 유지하고, 업무자료와 AI 문서 분석 결과만 주입."""
    html = ONBOARDING_SCHEDULE_HTML
    onboarding_data = _df_to_onboarding_data(schedule_df, project_df, contact_df, asset_df)
    replacements = {
        "EMBEDDED_SCHEDULE_TASKS": onboarding_data["schedule_tasks"],
        "EMBEDDED_PROJECT_DEADLINE_TASKS": onboarding_data["project_deadline_tasks"],
        "EMBEDDED_ASSET_DEADLINE_TASKS": onboarding_data["asset_deadline_tasks"],
        "EMBEDDED_PROJECT_OVERVIEW": onboarding_data["project_overview"],
        "EMBEDDED_ASSET_OVERVIEW": onboarding_data["asset_overview"],
        "EMBEDDED_CONTACTS": onboarding_data["contacts"],
    }
    for constant_name, value in replacements.items():
        html = _replace_js_array(html, constant_name, value)

    if ai_items:
        html = _replace_js_const(html, "AI_DOC_MODE", True)
        html = _replace_js_const(html, "AI_START_DATE", ai_items.get("start_date", ""))
        html = _replace_js_const(html, "AI_TIMED_EVENTS", ai_items.get("timed_events", []))
        html = _replace_js_const(html, "AI_ALLDAY_TASKS", ai_items.get("all_day_items", []))
        html = _replace_js_const(html, "AI_MAIL_REFS", ai_items.get("mails", []))

    uploaded_count = sum(df is not None for df in [schedule_df, project_df, contact_df, asset_df])
    if uploaded_count:
        html = html.replace(
            "업무일정·프로젝트 현황·자산·연락망 내용은 도구 안에 이미 반영되어 있어서 따로 업로드하지 않아도 됩니다.",
            f"📂 Excel 업무자료 탭에서 올린 자료 {uploaded_count}/4종이 자동으로 연동되어 있습니다.",
            1,
        )
    return html


# =========================================================
# 6. 후임자 대시보드
# =========================================================
def _to_number(value):
    try:
        return float(value)
    except Exception:
        return 0.0


def _to_datetime(value):
    if value is None:
        return pd.NaT
    if isinstance(value, pd.Timestamp):
        return value
    try:
        if pd.isna(value):
            return pd.NaT
    except Exception:
        pass
    if isinstance(value, (int, float)) and 30000 <= float(value) <= 70000:
        try:
            return pd.to_datetime(float(value), unit="D", origin="1899-12-30")
        except Exception:
            return pd.NaT
    return pd.to_datetime(value, errors="coerce")


def _format_dashboard_date(value):
    dt = _to_datetime(value)
    if pd.isna(dt):
        return safe_value(value)
    return dt.strftime("%Y-%m-%d")


def build_dashboard_summary(schedule_df, project_df, asset_df):
    summary = {"urgent_count": 0, "open_work_count": 0, "issue_count": 0, "pending_asset_count": 0,
               "average_progress": 0, "nearest_deadline": "", "nearest_work": "", "nearest_project": ""}

    if schedule_df is not None and not schedule_df.empty:
        summary["urgent_count"] = int(schedule_df["우선순위"].astype(str).str.strip().eq("긴급").sum())
        status = schedule_df["상태"].astype(str).str.strip()
        summary["open_work_count"] = int((~status.isin({"완료", "종료", "제출완료"}) & status.ne("")).sum())
        progress = pd.to_numeric(schedule_df["진행률"], errors="coerce")
        if progress.notna().any():
            summary["average_progress"] = int(round(progress.mean()))

        deadlines = schedule_df.copy()
        deadlines["_deadline"] = deadlines["마감"].apply(_to_datetime)
        deadlines = deadlines[deadlines["_deadline"].notna()]
        if not deadlines.empty:
            today = pd.Timestamp(date.today())
            upcoming = deadlines[deadlines["_deadline"].dt.normalize() >= today]
            if not upcoming.empty:
                nearest = upcoming.sort_values("_deadline").iloc[0]
            else:
                nearest = deadlines.sort_values("_deadline", ascending=False).iloc[0]
            summary["nearest_deadline"] = _format_dashboard_date(nearest.get("마감", ""))
            summary["nearest_work"] = safe_value(nearest.get("업무", ""))
            summary["nearest_project"] = safe_value(nearest.get("프로젝트/현장", ""))

    if project_df is not None and not project_df.empty:
        issues = project_df["리스크/이슈"].fillna("").astype(str).str.strip()
        summary["issue_count"] = int(issues.ne("").sum())

    if asset_df is not None and not asset_df.empty:
        status = asset_df["상태"].astype(str).str.strip()
        summary["pending_asset_count"] = int((~status.eq("완료") & status.ne("")).sum())

    return summary


def show_handover_dashboard(schedule_df, project_df, contact_df, asset_df):
    uploaded = [("업무일정", schedule_df), ("프로젝트", project_df), ("담당자", contact_df), ("권한·자산", asset_df)]
    uploaded_count = sum(df is not None for _, df in uploaded)

    st.header("🏠 업무 현황 대시보드")
    st.markdown(
        """
        <div class="section-note">
        후임자가 인수 직후 확인해야 할 <b>우선 업무, 마감, 프로젝트 리스크,
        계정·권한·자산 인계 상태</b>를 한 화면에서 확인합니다.
        </div>
        """,
        unsafe_allow_html=True,
    )

    if uploaded_count == 0:
        st.info("먼저 **📂 Excel 업무자료** 탭에서 업무자료를 업로드해주세요. 업로드하면 이 대시보드가 자동으로 채워집니다.")
        cols = st.columns(4)
        for column, (label, df) in zip(cols, uploaded):
            column.metric(label, "➖ 대기")
        return

    if uploaded_count < 4:
        missing = ", ".join(label for label, df in uploaded if df is None)
        st.caption(f"ℹ️ 아직 인식되지 않은 자료: {missing} — 있는 자료로만 대시보드를 표시합니다.")

    summary = build_dashboard_summary(schedule_df, project_df, asset_df)
    c1, c2, c3, c4, c5 = st.columns(5)
    c1.metric("🚨 긴급 업무", f"{summary['urgent_count']}건")
    c2.metric("🟡 진행/예정 업무", f"{summary['open_work_count']}건")
    c3.metric("⚠️ 미해결 이슈", f"{summary['issue_count']}건")
    c4.metric("🔐 미완료 인계", f"{summary['pending_asset_count']}건")
    c5.metric("📈 평균 진행률", f"{summary['average_progress']}%")

    st.divider()

    if schedule_df is not None and not schedule_df.empty:
        st.subheader("⏰ 가장 가까운 마감")
        if summary["nearest_deadline"]:
            st.warning(
                f"**{summary['nearest_deadline']}**까지 **{summary['nearest_project']} - "
                f"{summary['nearest_work']}** 업무를 확인해야 합니다."
            )
        else:
            st.info("등록된 마감 일정이 없습니다.")

        st.subheader("📌 후임자가 먼저 확인할 업무")
        view = schedule_df.copy()
        view["_deadline"] = view["마감"].apply(_to_datetime)
        view["_priority"] = view["우선순위"].astype(str).str.strip().map(
            {"긴급": 4, "상": 3, "중": 2, "하": 1}).fillna(0)
        view = view.sort_values(["_deadline", "_priority"], ascending=[True, False]).head(3)
        rows = []
        for _, row in view.iterrows():
            rows.append({
                "마감": _format_dashboard_date(row.get("마감", "")),
                "업무": safe_value(row.get("업무", "")),
                "프로젝트/현장": safe_value(row.get("프로젝트/현장", "")),
                "우선순위": safe_value(row.get("우선순위", "")),
                "진행률": f"{int(_to_number(row.get('진행률', 0)))}%",
                "다음 조치": safe_value(row.get("다음 조치", "")),
            })
        st.dataframe(pd.DataFrame(rows), use_container_width=True, hide_index=True)

    if project_df is not None and not project_df.empty:
        st.subheader("📊 프로젝트 진행 현황")
        columns = [c for c in ["프로젝트/현장", "현재 단계", "진행률", "다음 액션", "리스크/이슈", "마감일"]
                   if c in project_df.columns]
        project_view = project_df[columns].copy()
        if "진행률" in project_view.columns:
            project_view["진행률"] = project_view["진행률"].apply(lambda v: f"{int(_to_number(v))}%")
        if "마감일" in project_view.columns:
            project_view["마감일"] = project_view["마감일"].apply(_format_dashboard_date)
        st.dataframe(project_view, use_container_width=True, hide_index=True)

    left, right = st.columns(2)
    with left:
        st.subheader("⚠️ 미해결 이슈")
        issues = []
        if project_df is not None:
            for _, row in project_df.iterrows():
                issue = safe_value(row.get("리스크/이슈", ""))
                if issue:
                    issues.append((safe_value(row.get("프로젝트/현장", "")), issue))
        if issues:
            for project, issue in issues:
                st.write(f"• **{project}** — {issue}")
        else:
            st.success("등록된 미해결 이슈가 없습니다.")

    with right:
        st.subheader("🔐 남아 있는 인계")
        if asset_df is not None:
            status = asset_df["상태"].astype(str).str.strip()
            pending = asset_df[~status.eq("완료")]
            if not pending.empty:
                for _, row in pending.iterrows():
                    st.write(
                        f"• **{safe_value(row.get('시스템/자산', ''))}** — {safe_value(row.get('상태', ''))} "
                        f"/ 목표 {_format_dashboard_date(row.get('완료 목표일', ''))}"
                    )
            else:
                st.success("계정·권한·자산 인계가 모두 완료되었습니다.")
        else:
            st.caption("계정·권한·자산 자료가 없습니다.")

    st.divider()
    st.success(
        "💡 위 내용을 확인한 뒤 **💬 후임자 Q&A** 탭에서 세부 업무를 질문하거나, "
        "**🤖 업무메모 자동 인수인계**에서 최종 인수인계서를 생성할 수 있습니다."
    )


# =========================================================
# 7. 메인 화면 탭
# =========================================================
schedule_df = None
project_df = None
contact_df = None
asset_df = None

TABS_STATE_TRACKING_AVAILABLE = (
    "on_change" in inspect.signature(st.tabs).parameters
    and "key" in inspect.signature(st.tabs).parameters
)

if TABS_STATE_TRACKING_AVAILABLE:
    (tab_dashboard, tab_excel, tab_qa, tab_onboarding_schedule, tab_memo, tab_manual, tab_error_admin) = st.tabs(
        MAIN_SCREEN_OPTIONS, key="main_tabs", on_change="rerun")
else:
    (tab_dashboard, tab_excel, tab_qa, tab_onboarding_schedule, tab_memo, tab_manual, tab_error_admin) = st.tabs(
        MAIN_SCREEN_OPTIONS)


# =========================================================
# TAB. Excel 업무자료 업로드
# =========================================================
with tab_excel:
    st.header("📂 업무자료 통합 업로드")
    st.markdown(
        """
        <div class="section-note">
        업무일정·프로젝트 진행현황·담당자 연락망·계정/권한/자산 자료를 올리면 <b>AI가 파일 이름이나 열 이름이 달라도
        내용을 보고 어떤 자료인지 자동으로 인식</b>하고, 대시보드·Q&amp;A·온보딩 일정에 같은 데이터를 연결합니다.
        </div>
        """,
        unsafe_allow_html=True,
    )

    uploaded_excels = st.file_uploader(
        "업무자료(Excel / CSV)를 업로드하세요.",
        type=["xlsx", "csv"],
        accept_multiple_files=True,
        key="excel_upload",
    )

    if uploaded_excels:
        uploaded_names = [f.name for f in uploaded_excels]
        if uploaded_names != st.session_state.get("last_uploaded_files", []):
            track_action(screen="📂 Excel 업무자료", function="업무자료 업로드", action="Excel 파일 업로드",
                         input_value=", ".join(uploaded_names), metadata={"file_count": len(uploaded_names)})
            st.session_state["last_uploaded_files"] = uploaded_names

        files = [(f.name, f.getvalue()) for f in uploaded_excels]
        files_hash = content_hash(*[n for n, _ in files], *[b for _, b in files])

        with st.spinner("🤖 AI가 파일 구조를 분석하는 중이에요... (처음 한 번만)"):
            normalized = ai_try(
                "excel", files_hash, lambda: ai.normalize_excel_files(files),
                screen="📂 Excel 업무자료", function="AI 파일 구조 인식",
                code="AI_EXCEL_MAPPING_ERROR", input_value=", ".join(uploaded_names),
            )
        engine = "ai"
        if normalized is None:
            engine = "rule"
            for f in uploaded_excels:
                f.seek(0)
            normalized = rule_normalize_excel(uploaded_excels)

        schedule_df = normalized["schedule"]
        project_df = normalized["project"]
        contact_df = normalized["contact"]
        asset_df = normalized["asset"]
        unknown_files = normalized["unknown"]

        st.subheader("✅ 자료 자동 분류 결과")
        st.caption("🤖 AI가 내용을 보고 분류했습니다." if engine == "ai"
                   else "⚙️ 기존 규칙(파일 이름·고정 열 이름)으로 분류했습니다.")
        c1, c2, c3, c4 = st.columns(4)
        c1.metric("업무일정", "업로드 완료" if schedule_df is not None else "미업로드")
        c2.metric("프로젝트 현황", "업로드 완료" if project_df is not None else "미업로드")
        c3.metric("담당자 연락망", "업로드 완료" if contact_df is not None else "미업로드")
        c4.metric("계정·권한·자산", "업로드 완료" if asset_df is not None else "미업로드")

        if unknown_files:
            st.warning("자동 분류하지 못한 파일: " + ", ".join(unknown_files))

        if normalized.get("mapping"):
            with st.expander("🔎 AI가 인식한 표 구조 보기"):
                st.dataframe(pd.DataFrame(normalized["mapping"]), use_container_width=True, hide_index=True)

        st.divider()
        st.subheader("📋 인식된 업무자료 미리보기")
        for label, df in [("업무일정", schedule_df), ("프로젝트 현황", project_df),
                          ("담당자 연락망", contact_df), ("계정·권한·자산", asset_df)]:
            if df is None:
                continue
            with st.expander(f"📄 {label} · {df.attrs.get('source_file', '')}", expanded=False):
                st.caption(f"{len(df)}행 × {len(df.columns)}열")
                st.dataframe(df, use_container_width=True)

        st.success("업무자료 인식이 완료되었습니다. 대시보드·Q&A·온보딩 일정·인수인계서에 자동으로 연결됩니다.")
    else:
        st.caption("예시 파일: 01_업무일정.xlsx / 02_프로젝트_진행현황.xlsx / 03_담당자_연락망.xlsx / "
                   "04_계정_권한_자산.xlsx — 파일 이름이나 양식이 달라도 괜찮아요.")


# =========================================================
# TAB. 후임자 대시보드
# =========================================================
with tab_dashboard:
    try:
        show_handover_dashboard(schedule_df, project_df, contact_df, asset_df)
    except Exception as e:
        record_error(screen="🏠 대시보드", function="후임자 업무 대시보드", error=e, code="DASHBOARD_ERROR")
        st.error("대시보드를 그리는 중 오류가 발생했습니다. 시스템이 자동 감지했습니다.")


# =========================================================
# TAB. 후임자 Q&A
# =========================================================
with tab_qa:
    st.header("💬 근거 기반 업무 Q&A")
    st.markdown(
        """
        <div class="section-note">
        등록된 업무자료와 문서(회의록·메일)에서 질문과 관련된 정보를 찾아 <b>AI가 문맥을 이해해 답변</b>하고
        <b>근거 파일·시트·행</b>을 함께 표시합니다. 근거 위치는 AI가 아니라 시스템이 원본에서 직접 계산합니다.
        </div>
        """,
        unsafe_allow_html=True,
    )

    qa_items = [("업무일정", schedule_df), ("프로젝트", project_df), ("담당자", contact_df), ("권한·자산", asset_df)]
    uploaded_count = sum(df is not None for _, df in qa_items)
    shared_docs = get_shared_docs()

    if uploaded_count == 0 and not shared_docs:
        st.warning("먼저 📂 Excel 업무자료 탭에서 업무자료를 올리거나, 📅 온보딩 일과표 탭에서 문서를 올려주세요.")
    else:
        cols = st.columns(4)
        for column, (label, df) in zip(cols, qa_items):
            if df is not None:
                column.success(f"✅ {label}")
            else:
                column.info(f"➖ {label}")
        if shared_docs:
            st.caption(f"📎 회의록·캘린더·메일 문서 {len(shared_docs)}건도 함께 검색합니다.")

        reference_date = st.date_input(
            "기준일 (오늘로 간주할 날짜)",
            value=ai.default_reference_date(schedule_df),
            key="qa_reference_date",
            help="'이번 주', '오늘' 같은 질문을 이 날짜 기준으로 해석합니다. 기본값은 업무일정의 첫 날짜예요.",
        )

        st.markdown("#### 질문 예시")
        sample_questions = [
            "직접 입력",
            "9월 1일에 가장 먼저 해야 할 업무가 뭐야?",
            "A동 보고서는 지금 어디까지 진행됐어?",
            "A동 보고서가 늦어질 수 있는 이유가 있어?",
            "B공장 견적이 아직 완료되지 않은 이유는 뭐야?",
            "B공장 견적 다음 액션은?",
            "C센터 회의 전에 뭘 준비해야 해?",
            "C센터에서 아직 결정 안 된 게 뭐야?",
            "A동 고객사 담당자는 누구야?",
            "박현우 부장님께 보고서 보낼 때 주의할 점 있어?",
            "감지기 단가 문의는 누구한테 해야 해?",
            "후임자가 공용드라이브를 바로 쓸 수 있어?",
            "비밀번호를 인계자한테 받아야 하는 시스템이 있어?",
            "이번 주 긴급한 업무만 정리해줘.",
            "9월 3일 일정 알려줘.",
            "진행률이 가장 높은 프로젝트는 뭐야?",
            "진행률 50% 이상인 업무를 알려줘.",
            "외부 담당자 중 이메일로 먼저 연락할 사람이 누구야?",
            "후임자가 첫 주에 권한 관련해서 해야 할 일은?",
            "지금 미해결 이슈가 있는 프로젝트를 정리해줘.",
            "후임자가 가장 먼저 확인할 3가지만 뽑아줘.",
            "고객사에 자료 보내기 전에 내부 확인이 필요한 업무가 있어?",
            "9월 4일까지 끝내야 할 인수인계 항목이 뭐야?",
            "B공장 관련해서 누구랑 협업해야 해?",
            "신규로 인계받은 뒤 첫 신규 현장은 어디야?",
            "완료된 권한 인계와 아직 남은 권한 인계를 구분해줘.",
            "C센터 미조치 2건 일정은 언제까지 확정해야 해?",
        ]
        selected_question = st.selectbox("예시 질문 선택", sample_questions, key="qa_sample_question")
        typed_question = st.text_input("후임자가 궁금한 내용을 입력하세요",
                                       placeholder="예: A동 보고서는 지금 어디까지 진행됐어?", key="qa_question")

        if typed_question.strip():
            final_question = typed_question.strip()
        elif selected_question != "직접 입력":
            final_question = selected_question
        else:
            final_question = ""

        if st.button("🔎 질문하기", type="primary", use_container_width=True, key="qa_ask_button"):
            if not final_question:
                st.warning("질문을 입력하거나 예시 질문을 선택해주세요.")
            else:
                with st.chat_message("user"):
                    st.write(final_question)
                track_action(screen="💬 후임자 Q&A", function="후임자 질문 검색",
                             action="질문하기 버튼 클릭", input_value=final_question)

                qa_started_at = time.perf_counter()
                result = None
                try:
                    with st.spinner("🤖 자료를 찾아보는 중이에요..."):
                        result = ai.answer_question(
                            final_question,
                            schedule_df=schedule_df, project_df=project_df,
                            contact_df=contact_df, asset_df=asset_df,
                            docs=shared_docs, reference_date=reference_date.isoformat(),
                        )
                    qa_elapsed = time.perf_counter() - qa_started_at
                    st.session_state["last_response_time"] = qa_elapsed
                    response_by_screen = dict(st.session_state.get("last_response_time_by_screen", {}))
                    response_by_screen["💬 후임자 Q&A"] = qa_elapsed
                    st.session_state["last_response_time_by_screen"] = response_by_screen
                    track_action(screen="💬 후임자 Q&A", function="후임자 질문 검색", action="답변 생성 완료",
                                 input_value=final_question,
                                 metadata={"response_time": round(qa_elapsed, 3), "engine": result.get("engine")})

                    if result.get("engine") == "rule" and result.get("fallback_reason") and ai.ai_available():
                        record_system_issue(
                            screen="💬 후임자 Q&A", function="후임자 질문 검색", code="AI_QA_FALLBACK",
                            message=f"AI 답변 실패로 규칙 기반 답변으로 전환: {result['fallback_reason']}",
                            issue_type="AI_FALLBACK", input_value=final_question, response_time=qa_elapsed,
                        )
                    elif qa_elapsed >= SLOW_RESPONSE_SECONDS:
                        record_system_issue(
                            screen="💬 후임자 Q&A", function="후임자 질문 검색", code="QA_SLOW_RESPONSE",
                            message=f"Q&A 응답 시간이 기준({SLOW_RESPONSE_SECONDS:.0f}초)을 초과했습니다: {qa_elapsed:.2f}초",
                            issue_type="SLOW_RESPONSE", input_value=final_question, response_time=qa_elapsed,
                        )
                except Exception as e:
                    qa_elapsed = time.perf_counter() - qa_started_at
                    record_error(screen="💬 후임자 Q&A", function="후임자 질문 검색", error=e,
                                 code="QA_ANSWER_ERROR", input_value=final_question, response_time=qa_elapsed)
                    st.error("❌ 답변 생성 중 오류가 발생했습니다. 시스템이 자동 감지했으며 오류 리포트 창이 열립니다.")

                if result is not None:
                    with st.chat_message("assistant"):
                        st.markdown(result["answer"])
                        if result.get("engine") == "ai":
                            st.caption("🤖 AI 답변")
                        elif result.get("engine") == "rule":
                            st.caption("⚙️ 규칙 기반 답변" + (" (AI 연결 실패로 전환)" if result.get("fallback_reason") else ""))

                        if result.get("sources"):
                            st.markdown("##### 📎 답변 근거")
                            for source in result["sources"]:
                                if source.get("row"):
                                    location = f"{source['sheet']} {source['row']}행"
                                else:
                                    location = source.get("sheet") or "문서"
                                st.caption(f"• {source['file']} · {location}")
                                if source.get("evidence"):
                                    st.caption(f"  ↳ {source['evidence']}")
                        else:
                            st.caption("연결된 근거 자료가 없습니다.")

        st.divider()
        with st.expander("ℹ️ 현재 Q&A 방식"):
            st.write(
                "AI가 업로드된 표의 모든 행과 문서를 읽고 질문의 뜻을 이해해 답변합니다. "
                "표현이 달라도, 여러 파일을 연결해야 하는 질문도 답할 수 있습니다."
            )
            st.write(
                "각 행과 문서에는 번호표(ID)가 붙어 있고, AI는 사용한 번호표만 알려줍니다. "
                "파일명·행 번호는 시스템이 원본에서 계산하므로 근거가 지어내질 수 없습니다."
            )
            st.write("AI를 쓸 수 없을 때는 기존 규칙 기반 검색으로 자동 전환됩니다.")


# =========================================================
# TAB. 후임자 온보딩 일과표
# =========================================================
with tab_onboarding_schedule:
    st.header("📅 후임자 온보딩 일정 설계")
    st.markdown(
        """
        <div class="section-note">
        Excel 업무자료와 회의록·캘린더·메일 문서를 합쳐 <b>1주차·2주차 온보딩 일과표</b>를 생성합니다.
        AI가 문서 형식과 상관없이 회의 일정, 마감, 회의록의 후속 조치(Next Action), 메일 내용을 뽑아냅니다.
        생성 후 신입/선임자 역할을 선택해 일정 칸을 직접 수정할 수 있습니다.
        </div>
        """,
        unsafe_allow_html=True,
    )

    connected_items = [("업무일정", schedule_df), ("프로젝트", project_df), ("담당자", contact_df), ("권한·자산", asset_df)]
    connected_count = sum(df is not None for _, df in connected_items)
    if connected_count:
        st.success(f"🔗 **Excel 업무자료 {connected_count}/4종 연동 중** · 같은 자료를 다시 올릴 필요가 없습니다.")
    else:
        st.info("아직 📂 Excel 업무자료를 올리지 않아 데모 데이터로 일정표를 만듭니다.")

    st.markdown("#### 📎 회의록 · 캘린더 · 메일 문서")
    st.file_uploader(
        "문서를 한 번에 모두 올려주세요 (docx / pdf / txt, 여러 개 가능). 여기 올린 문서는 Q&A와 인수인계서 생성에도 쓰입니다.",
        type=["docx", "pdf", "txt", "md"],
        accept_multiple_files=True,
        key="onboarding_docs",
    )
    onboarding_docs = get_shared_docs()

    ai_items = None
    if onboarding_docs:
        docs_hash = content_hash(*[n for n, _ in onboarding_docs], *[t for _, t in onboarding_docs])
        if ai.ai_available():
            with st.spinner("🤖 AI가 문서에서 일정을 찾는 중이에요... (처음 한 번만, 20~60초)"):
                ai_items = ai_try(
                    "schedule", docs_hash, lambda: ai.extract_schedule_items(onboarding_docs),
                    screen="📅 온보딩 일과표", function="AI 문서 일정 추출",
                    code="AI_SCHEDULE_EXTRACT_ERROR", input_value=f"문서 {len(onboarding_docs)}건",
                )
            if ai_items:
                c1, c2, c3, c4 = st.columns(4)
                c1.metric("시간 일정", f"{len(ai_items['timed_events'])}건")
                c2.metric("마감·조치", f"{len(ai_items['all_day_items'])}건")
                c3.metric("메일", f"{len(ai_items['mails'])}건")
                c4.metric("인수 시작일", ai_items["start_date"] or "자동")
                with st.expander("🔎 AI가 문서에서 찾은 내용 보기"):
                    if ai_items["timed_events"]:
                        st.markdown("**시간 일정**")
                        st.dataframe(pd.DataFrame(ai_items["timed_events"]), use_container_width=True, hide_index=True)
                    if ai_items["all_day_items"]:
                        st.markdown("**마감 · 조치**")
                        st.dataframe(pd.DataFrame(ai_items["all_day_items"]), use_container_width=True, hide_index=True)
                    if ai_items["mails"]:
                        st.markdown("**메일**")
                        st.dataframe(pd.DataFrame(ai_items["mails"]), use_container_width=True, hide_index=True)
            elif st.session_state.get(f"_ai_failed_schedule_{docs_hash}"):
                st.warning("AI 문서 분석에 실패했어요. 아래 일정표 왼쪽의 업로드 칸을 이용하면 기존 방식으로 만들 수 있어요.")
                if st.button("🔄 AI 분석 다시 시도", key="retry_schedule_ai"):
                    st.session_state.pop(f"_ai_failed_schedule_{docs_hash}", None)
                    st.rerun()
        else:
            st.info("AI가 연결되지 않아 문서 분석을 할 수 없어요. 아래 일정표 왼쪽의 업로드 칸에 docx를 올리면 기존 방식으로 만들 수 있어요.")

    try:
        onboarding_html = build_onboarding_schedule_html(schedule_df, project_df, contact_df, asset_df, ai_items=ai_items)
        components.html(onboarding_html, height=1500, scrolling=True)
    except Exception as e:
        record_error(screen="📅 온보딩 일과표", function="Excel 업무자료 일정표 연동", error=e,
                     code="ONBOARDING_DATA_LINK_ERROR")
        st.error("❌ 업무자료를 일정표에 연결하는 중 오류가 발생했습니다. 오류가 자동 기록되었습니다.")


# =========================================================
# TAB. 인수인계서 자동 생성 (업무메모 / 전체 자료)
# =========================================================
TASK_COLUMNS = ["주요 업무", "업무 목적", "업무 프로세스", "우선순위", "비고"]
DETAIL_FIELDS = [
    "업무 개요", "목적 / 성과지표", "진행 중 프로젝트 현황", "정기 업무", "비정기 업무",
    "주요 일정 / 마감", "관련 시스템 / 계정 / 권한", "관련 담당자 / 연락처", "협업 부서",
    "특이사항 / 주의사항", "리스크 / 미해결 이슈", "참고 파일 경로 / 문서 링크", "후임자 숙지 필요사항",
]
URGENT_COLUMNS = ["일자", "내용", "대응 방법", "담당"]
MONTHLY_COLUMNS = ["일자", "일정 내용", "비고"]
ASSET_COLUMNS = ["시스템 / 자산명", "유형", "권한 수준", "인계 방법", "상태", "비고"]
CHECKLIST_COLUMNS = ["체크", "항목", "확인일", "확인자", "비고"]


def reset_review_widgets_if_changed(parsed_data):
    """새 초안이 들어오면 이전 검토 화면 입력값을 비워서 새 내용이 보이게 한다."""
    data_hash = content_hash(json.dumps(parsed_data, ensure_ascii=False, sort_keys=True))
    if st.session_state.get("_review_data_hash") != data_hash:
        for key in list(st.session_state.keys()):
            if key.startswith("auto_") and key != "auto_memo_upload":
                del st.session_state[key]
        st.session_state["_review_data_hash"] = data_hash


def render_review_and_download(parsed_data, source_texts):
    """AI 또는 규칙 기반으로 만든 초안을 검토·수정하고 DOCX로 내려받는 화면 (기존과 동일)."""
    with st.expander("📋 원본 자료 보기"):
        for name, text in source_texts:
            st.markdown(f"**{name}**")
            st.text(text[:5000] + ("\n...(이하 생략)" if len(text) > 5000 else ""))

    st.subheader("Ⅰ. 기본 정보")
    c1, c2 = st.columns(2)
    with c1:
        auto_institution = st.text_input("기관명", value=parsed_data["meta"].get("기관명", ""), key="auto_institution")
        auto_department_meta = st.text_input("부서명", value=parsed_data["meta"].get("부서명", ""), key="auto_department_meta")
        auto_department = st.text_input("소속 부서", value=parsed_data["basic"].get("소속 부서", ""), key="auto_department")
        auto_giver = st.text_input("인계자 성명", value=parsed_data["basic"].get("인계자 성명", ""), key="auto_giver")
        auto_written_date = st.text_input("작성일", value=parsed_data["basic"].get("작성일", ""), key="auto_written_date")
    with c2:
        auto_document_no = st.text_input("문서번호", value=parsed_data["meta"].get("문서번호", ""), key="auto_document_no")
        auto_retention = st.text_input("보존기간", value=parsed_data["meta"].get("보존기간", ""), key="auto_retention")
        auto_position = st.text_input("직위 / 직책", value=parsed_data["basic"].get("직위 / 직책", ""), key="auto_position")
        auto_receiver = st.text_input("인수자 성명", value=parsed_data["basic"].get("인수자 성명", ""), key="auto_receiver")
        auto_expected_date = st.text_input("인수인계 완료 예정일",
                                           value=parsed_data["basic"].get("인수인계 완료 예정일", ""), key="auto_expected_date")

    st.subheader("Ⅱ. 담당업무 개요")
    auto_tasks_df = st.data_editor(pd.DataFrame(parsed_data.get("tasks", []), columns=TASK_COLUMNS),
                                   num_rows="dynamic", use_container_width=True, key="auto_tasks")

    st.subheader("Ⅲ. 업무 상세")
    edited_details = []
    for i, detail in enumerate(parsed_data.get("details", [])):
        task_name = detail.get("업무명", f"업무 {i + 1}")
        with st.expander(f"📌 {i + 1}. {task_name}", expanded=(i == 0)):
            edited_detail = {"업무명": st.text_input("업무명", value=task_name, key=f"auto_detail_name_{i}")}
            for field in DETAIL_FIELDS:
                edited_detail[field] = st.text_area(field, value=detail.get(field, ""), key=f"auto_detail_{i}_{field}")
            status_options = ["미완료", "진행중", "완료"]
            current_status = detail.get("인수인계 완료 여부", "미완료")
            if current_status not in status_options:
                current_status = "미완료"
            edited_detail["인수인계 완료 여부"] = st.selectbox(
                "인수인계 완료 여부", status_options, index=status_options.index(current_status), key=f"auto_done_{i}")
            edited_details.append(edited_detail)

    st.subheader("Ⅳ. 주요 일정 및 대외 커뮤니케이션")
    st.markdown("#### 1. 긴급 업무 일정")
    auto_urgent_df = st.data_editor(pd.DataFrame(parsed_data.get("urgent_schedule", []), columns=URGENT_COLUMNS),
                                    num_rows="dynamic", use_container_width=True, key="auto_urgent")
    st.markdown("#### 2. 향후 1개월 주요 일정")
    auto_monthly_df = st.data_editor(pd.DataFrame(parsed_data.get("monthly_schedule", []), columns=MONTHLY_COLUMNS),
                                     num_rows="dynamic", use_container_width=True, key="auto_monthly")
    st.markdown("#### 3. 대외 커뮤니케이션 유의사항")
    auto_communication = st.text_area("커뮤니케이션 유의사항", value=parsed_data.get("communication_note", ""),
                                      key="auto_communication")

    st.subheader("Ⅴ. 계정 · 권한 · 자산 인계")
    auto_assets_df = st.data_editor(pd.DataFrame(parsed_data.get("assets", []), columns=ASSET_COLUMNS),
                                    num_rows="dynamic", use_container_width=True, key="auto_assets")

    st.subheader("Ⅵ. 인수인계 체크리스트")
    auto_checklist_df = st.data_editor(pd.DataFrame(parsed_data.get("checklist", []), columns=CHECKLIST_COLUMNS),
                                       num_rows="dynamic", use_container_width=True, key="auto_checklist")

    st.subheader("Ⅶ. 서명")
    signatures = parsed_data.get("signatures", {})
    s1, s2, s3 = st.columns(3)
    with s1:
        auto_sign_giver = st.text_input("인계자", value=signatures.get("인계자 성명", auto_giver), key="auto_sign_giver")
        auto_sign_giver_date = st.text_input("인계자 서명일", value=signatures.get("인계자 서명일", ""), key="auto_sign_giver_date")
    with s2:
        auto_sign_receiver = st.text_input("인수자", value=signatures.get("인수자 성명", auto_receiver), key="auto_sign_receiver")
        auto_sign_receiver_date = st.text_input("인수자 서명일", value=signatures.get("인수자 서명일", ""), key="auto_sign_receiver_date")
    with s3:
        auto_sign_checker = st.text_input("확인자(팀장 등)", value=signatures.get("확인자 성명", ""), key="auto_sign_checker")
        auto_sign_checker_date = st.text_input("확인자 서명일", value=signatures.get("확인자 서명일", ""), key="auto_sign_checker_date")

    reviewed_data = {
        "meta": {"기관명": auto_institution, "부서명": auto_department_meta,
                 "문서번호": auto_document_no, "보존기간": auto_retention},
        "basic": {"소속 부서": auto_department, "직위 / 직책": auto_position, "인계자 성명": auto_giver,
                  "인수자 성명": auto_receiver, "작성일": auto_written_date, "인수인계 완료 예정일": auto_expected_date},
        "tasks": clean_records(auto_tasks_df),
        "details": edited_details,
        "urgent_schedule": clean_records(auto_urgent_df),
        "monthly_schedule": clean_records(auto_monthly_df),
        "communication_note": auto_communication,
        "assets": clean_records(auto_assets_df),
        "checklist": clean_records(auto_checklist_df),
        "signatures": {"인계자 성명": auto_sign_giver, "인계자 서명일": auto_sign_giver_date,
                       "인수자 성명": auto_sign_receiver, "인수자 서명일": auto_sign_receiver_date,
                       "확인자 성명": auto_sign_checker, "확인자 서명일": auto_sign_checker_date},
    }

    st.divider()
    show_readiness(reviewed_data)
    st.divider()

    st.subheader("📄 최종 인수인계서 생성")
    st.info("위 내용을 확인하거나 수정한 뒤 아래 버튼으로 최종 문서를 생성하세요.")
    st.download_button(
        "✅ 검토 완료 - 최종 인수인계서 다운로드",
        data=build_docx(reviewed_data),
        file_name="업무_인수인계서_최종.docx",
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        use_container_width=True,
        key="auto_final_download",
    )
    with st.expander("🔍 최종 데이터 확인"):
        st.json(reviewed_data)


with tab_memo:
    st.header("🤖 인수인계서 자동 생성")
    st.markdown(
        """
        <div class="section-note">
        <b>업무메모 한 건</b>(형식 자유: txt·docx·pdf)을 올리거나, <b>지금까지 올린 업무자료와 문서 전체</b>를 AI가 종합해
        인수인계서 초안을 만듭니다. 사용자가 검토·수정한 뒤 준비도와 누락 항목을 확인하고 최종 DOCX를 생성합니다.
        </div>
        """,
        unsafe_allow_html=True,
    )

    memo_mode = st.radio(
        "만드는 방법",
        ["📄 업무메모 파일로 만들기", "📚 업로드한 전체 자료로 만들기 (AI)"],
        horizontal=True,
        key="memo_mode",
    )

    parsed_data = None
    source_texts = []

    if memo_mode.startswith("📄"):
        uploaded_memo = st.file_uploader("업무메모 파일을 업로드하세요 (txt / docx / pdf / md)",
                                         type=["txt", "docx", "pdf", "md"], key="auto_memo_upload")
        if uploaded_memo is not None:
            try:
                memo_text = ai.extract_text(uploaded_memo.name, uploaded_memo.getvalue())
                source_texts = [(uploaded_memo.name, memo_text)]
                memo_hash = content_hash(uploaded_memo.name, memo_text)

                with st.spinner("🤖 AI가 업무메모를 분석하는 중이에요... (처음 한 번만, 30~90초)"):
                    parsed_data = ai_try(
                        "memo", memo_hash, lambda: ai.generate_handover(source_texts, mode="memo"),
                        screen="🤖 업무메모 자동 인수인계", function="AI 업무메모 분석",
                        code="AI_MEMO_ERROR", input_value=uploaded_memo.name,
                    )
                if parsed_data is not None:
                    st.success("✅ AI가 업무메모 분석을 마쳤어요. 아래 내용을 확인하고 수정해주세요.")
                elif uploaded_memo.name.lower().endswith((".txt", ".md")):
                    parsed_data = parse_memo_text(memo_text)
                    st.success("✅ 업무메모 분석 완료 (⚙️ 규칙 기반). 아래 내용을 확인하고 수정해주세요.")
                    st.caption("규칙 기반 모드는 '키: 값'과 '|' 형식을 지킨 메모만 정확히 읽을 수 있어요.")
                else:
                    st.error("AI가 연결되지 않아 docx/pdf 메모는 분석할 수 없어요. txt 형식 메모를 올려주세요.")
            except Exception as e:
                record_error(screen="🤖 업무메모 자동 인수인계", function="업무메모 읽기", error=e,
                             code="MEMO_READ_ERROR", input_value=getattr(uploaded_memo, "name", "업무메모"))
                st.error("업무메모를 읽는 중 오류가 발생했습니다. 시스템이 자동 감지했습니다.")
        else:
            st.caption("업무메모_예시.txt 파일을 올려 테스트할 수 있어요. AI가 연결되어 있으면 형식이 달라도 괜찮아요.")

    else:
        excel_sources = get_excel_sources()
        doc_sources = get_shared_docs()
        extra_files = st.file_uploader(
            "추가로 넣을 자료가 있으면 올려주세요 (선택)",
            type=["xlsx", "csv", "docx", "pdf", "txt", "md"], accept_multiple_files=True, key="handover_extra",
        )
        extra_sources = []
        for f in extra_files or []:
            try:
                extra_sources.append((f.name, ai.extract_text(f.name, f.getvalue())))
            except Exception as e:
                st.warning(f"{f.name} 을(를) 읽지 못했어요: {e}")

        source_texts = excel_sources + doc_sources + extra_sources
        c1, c2, c3 = st.columns(3)
        c1.metric("Excel 업무자료", f"{len(excel_sources)}개")
        c2.metric("회의록·캘린더·메일", f"{len(doc_sources)}개")
        c3.metric("추가 자료", f"{len(extra_sources)}개")

        if not ai.ai_available():
            st.warning("이 방법은 AI가 연결되어 있어야 사용할 수 있어요. 위쪽 안내를 확인해주세요.")
        elif not source_texts:
            st.info("먼저 📂 Excel 업무자료 탭이나 📅 온보딩 일과표 탭에서 자료를 올려주세요.")
        else:
            sources_hash = content_hash(*[n for n, _ in source_texts], *[t for _, t in source_texts])
            if st.button("🤖 AI로 인수인계서 초안 만들기", type="primary", use_container_width=True,
                         key="handover_generate_button"):
                track_action(screen="🤖 업무메모 자동 인수인계", function="AI 인수인계서 종합 생성",
                             action="초안 만들기 버튼 클릭", input_value=f"자료 {len(source_texts)}개")
                started = time.perf_counter()
                try:
                    with st.spinner("🤖 AI가 자료 전체를 읽고 인수인계서를 쓰는 중이에요... (1~2분 걸릴 수 있어요)"):
                        generated = ai.generate_handover(source_texts, mode="integrated")
                    st.session_state["handover_generated"] = {"hash": sources_hash, "data": generated}
                except Exception as e:
                    record_error(screen="🤖 업무메모 자동 인수인계", function="AI 인수인계서 종합 생성", error=e,
                                 code="AI_HANDOVER_ERROR", input_value=f"자료 {len(source_texts)}개",
                                 response_time=time.perf_counter() - started)
                    st.error("AI 초안 생성 중 오류가 발생했습니다. 시스템이 자동 감지했습니다.")

            saved = st.session_state.get("handover_generated")
            if saved:
                parsed_data = saved["data"]
                if saved["hash"] != sources_hash:
                    st.info("자료가 바뀌었어요. 새 자료로 다시 만들려면 위 버튼을 눌러주세요. (아래는 이전 초안)")
                else:
                    st.success("✅ AI가 인수인계서 초안을 만들었어요. 아래 내용을 확인하고 수정해주세요.")

    if parsed_data is not None:
        try:
            reset_review_widgets_if_changed(parsed_data)
            render_review_and_download(parsed_data, source_texts)
        except Exception as e:
            record_error(screen="🤖 업무메모 자동 인수인계", function="인수인계서 검토 및 문서 생성", error=e,
                         code="MEMO_PROCESS_ERROR")
            st.error("인수인계서 처리 중 오류가 발생했습니다. 시스템이 자동 감지했습니다.")


# =========================================================
# TAB. 직접 작성
# =========================================================
with tab_manual:
    st.header("✍️ 인수인계서 직접 작성")
    st.caption("업무메모 없이도 직접 내용을 입력하여 DOCX를 생성할 수 있습니다.")

    with st.expander("문서 기본 정보", expanded=True):
        c1, c2, c3, c4 = st.columns(4)
        institution = c1.text_input("기관명", key="manual_institution")
        department_meta = c2.text_input("부서명", key="manual_department_meta")
        document_no = c3.text_input("문서번호", key="manual_document_no")
        retention = c4.text_input("보존기간", placeholder="예: 3년 / 폐기 시", key="manual_retention")

    st.subheader("I. 기본 정보")
    c1, c2 = st.columns(2)
    with c1:
        department = st.text_input("소속 부서", key="manual_department")
        giver = st.text_input("인계자 성명", key="manual_giver")
        written_date = st.date_input("작성일", value=date.today(), key="manual_written_date")
    with c2:
        position = st.text_input("직위 / 직책", key="manual_position")
        receiver = st.text_input("인수자 성명", key="manual_receiver")
        expected_date = st.date_input("인수인계 완료 예정일", value=date.today(), key="manual_expected_date")

    st.subheader("II. 담당업무 개요")
    tasks_df = st.data_editor(default_rows(TASK_COLUMNS, 3), num_rows="dynamic", use_container_width=True, key="manual_tasks")

    st.subheader("III. 업무 상세")
    st.caption("MVP에서는 최대 3개 주요 업무를 입력합니다.")
    details = []
    left_fields = ["업무 개요", "목적 / 성과지표", "진행 중 프로젝트 현황", "정기 업무", "비정기 업무",
                   "주요 일정 / 마감", "관련 시스템 / 계정 / 권한"]
    right_fields = ["관련 담당자 / 연락처", "협업 부서", "특이사항 / 주의사항", "리스크 / 미해결 이슈",
                    "참고 파일 경로 / 문서 링크", "후임자 숙지 필요사항"]
    for idx, task_tab in enumerate(st.tabs(["업무 1", "업무 2", "업무 3"]), start=1):
        with task_tab:
            task_name = st.text_input(f"업무명 {idx}", key=f"manual_task_name_{idx}")
            col1, col2 = st.columns(2)
            values = {}
            with col1:
                for field in left_fields:
                    values[field] = st.text_area(field, key=f"manual_{idx}_{field}")
            with col2:
                for field in right_fields:
                    values[field] = st.text_area(field, key=f"manual_{idx}_{field}")
                done = st.selectbox("인수인계 완료 여부", ["미완료", "진행중", "완료"], key=f"manual_done_{idx}")
            if task_name.strip() or any(v.strip() for v in values.values()):
                details.append({"업무명": task_name, **values, "인수인계 완료 여부": done})

    st.subheader("IV. 주요 일정 및 대외 커뮤니케이션")
    st.markdown("#### 1. 긴급 업무 일정")
    urgent_df = st.data_editor(default_rows(URGENT_COLUMNS, 2), num_rows="dynamic", use_container_width=True, key="manual_urgent")
    st.markdown("#### 2. 향후 1개월 주요 일정")
    monthly_df = st.data_editor(default_rows(MONTHLY_COLUMNS, 3), num_rows="dynamic", use_container_width=True, key="manual_monthly")
    st.markdown("#### 3. 대외 커뮤니케이션 유의사항")
    communication_note = st.text_area("커뮤니케이션 유의사항", key="manual_communication")

    st.subheader("V. 계정 · 권한 · 자산 인계")
    assets_df = st.data_editor(default_rows(ASSET_COLUMNS, 3), num_rows="dynamic", use_container_width=True, key="manual_assets")

    st.subheader("VI. 인수인계 체크리스트")
    check_default = pd.DataFrame([
        {"체크": "미완료", "항목": item, "확인일": "", "확인자": "", "비고": ""}
        for item in ["후임자 우선 숙지사항 전달", "참고 파일 경로 및 링크 전달",
                     "계정 및 권한 이전 필요사항 전달", "관련 담당자 및 연락처 전달"]
    ])
    checklist_df = st.data_editor(check_default, num_rows="dynamic", use_container_width=True, key="manual_checklist")

    st.subheader("VII. 서명")
    s1, s2, s3 = st.columns(3)
    with s1:
        sign_giver = st.text_input("인계자 성명", value=giver, key="manual_sign_giver")
        sign_giver_date = st.date_input("인계자 서명일", value=date.today(), key="manual_sign_giver_date")
    with s2:
        sign_receiver = st.text_input("인수자 성명", value=receiver, key="manual_sign_receiver")
        sign_receiver_date = st.date_input("인수자 서명일", value=date.today(), key="manual_sign_receiver_date")
    with s3:
        sign_checker = st.text_input("확인자(팀장 등) 성명", key="manual_sign_checker")
        sign_checker_date = st.date_input("확인자 서명일", value=date.today(), key="manual_sign_checker_date")

    manual_data = {
        "meta": {"기관명": institution, "부서명": department_meta, "문서번호": document_no, "보존기간": retention},
        "basic": {"소속 부서": department, "직위 / 직책": position, "인계자 성명": giver, "인수자 성명": receiver,
                  "작성일": written_date.isoformat(), "인수인계 완료 예정일": expected_date.isoformat()},
        "tasks": clean_records(tasks_df),
        "details": details,
        "urgent_schedule": clean_records(urgent_df),
        "monthly_schedule": clean_records(monthly_df),
        "communication_note": communication_note,
        "assets": clean_records(assets_df),
        "checklist": clean_records(checklist_df),
        "signatures": {"인계자 성명": sign_giver, "인계자 서명일": sign_giver_date.isoformat(),
                       "인수자 성명": sign_receiver, "인수자 서명일": sign_receiver_date.isoformat(),
                       "확인자 성명": sign_checker, "확인자 서명일": sign_checker_date.isoformat()},
    }

    st.divider()
    show_readiness(manual_data)
    st.divider()

    st.subheader("📄 자동 생성")
    c1, c2 = st.columns(2)
    with c1:
        st.download_button("💾 입력 데이터(JSON) 저장",
                           data=json.dumps(manual_data, ensure_ascii=False, indent=2).encode("utf-8"),
                           file_name="handover_data.json", mime="application/json",
                           use_container_width=True, key="manual_json_download")
    with c2:
        st.download_button("📄 인수인계서(DOCX) 생성", data=build_docx(manual_data),
                           file_name="업무_인수인계서_자동생성.docx",
                           mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                           use_container_width=True, key="manual_docx_download")
    with st.expander("현재 저장될 데이터 미리보기"):
        st.json(manual_data)


# =========================================================
# TAB. 관리자 오류함
# =========================================================
with tab_error_admin:
    st.header("🛠 관리자 오류함")
    st.caption(
        "시스템이 자동 감지한 오류·응답 지연 정보를 사용자가 확인 후 전송하면 이곳에 접수됩니다. "
        "현재 MVP에서는 현재 세션과 로컬 JSONL 파일에 기록되며, 배포 환경에서는 영구 저장소 연동으로 확장 가능합니다."
    )

    with st.expander("🧪 자동 오류 감지 테스트"):
        st.caption("시연용 버튼입니다. 누르면 시스템이 오류를 자동 감지한 것처럼 리포트 팝업이 자동으로 열립니다.")
        if st.button("⚡ 테스트 오류 발생시키기", key="make_demo_error"):
            try:
                raise TimeoutError("시연용: Q&A 응답 시간이 초과되었습니다.")
            except Exception as e:
                record_error(screen="💬 후임자 Q&A", function="후임자 질문 검색", error=e, code="QA_TIMEOUT_DEMO",
                             input_value="A동 보고서는 지금 어디까지 진행됐어?", response_time=31.24)
                st.info("테스트 오류를 자동 감지했습니다. 오류 리포트 창이 자동으로 열립니다.")

    error_reports = list(reversed(st.session_state.get("error_reports", [])))
    if not error_reports:
        st.info("아직 접수된 오류 신고가 없습니다.")
    else:
        report_rows = []
        for report in error_reports:
            response_time = report.get("response_time")
            report_rows.append({
                "신고번호": report.get("report_id", ""),
                "신고시각": report.get("reported_at", ""),
                "화면": report.get("screen", ""),
                "기능": report.get("function", ""),
                "에러 코드": report.get("error_code", ""),
                "응답시간(초)": round(float(response_time), 2) if response_time is not None else "",
                "상태": report.get("status", "미처리"),
            })
        st.dataframe(pd.DataFrame(report_rows), use_container_width=True, hide_index=True)

        selected_report_id = st.selectbox("상세 확인할 신고", [r.get("report_id", "") for r in error_reports],
                                          key="admin_error_report_select")
        selected_report = next((r for r in error_reports if r.get("report_id") == selected_report_id), None)

        if selected_report:
            action_col, delete_col = st.columns([3, 1])
            with action_col:
                st.caption(f"선택된 신고: {selected_report_id} · {selected_report.get('screen', '')} · "
                           f"{selected_report.get('reported_at', '')}")
            with delete_col:
                delete_confirmed = st.checkbox("삭제 확인", key=f"delete_confirm_{selected_report_id}",
                                               help="실수로 삭제하는 것을 막기 위한 확인 단계입니다.")
                if st.button("🗑 선택 신고 삭제", use_container_width=True, disabled=not delete_confirmed,
                             key=f"delete_error_report_{selected_report_id}"):
                    try:
                        if delete_error_report(selected_report_id):
                            st.success(f"{selected_report_id} 신고를 삭제했습니다.")
                            st.rerun()
                        else:
                            st.warning("선택한 신고를 찾지 못했습니다. 화면을 새로고침한 뒤 다시 확인해주세요.")
                    except Exception as e:
                        st.error(f"신고 삭제 중 오류가 발생했습니다: {e}")

            st.markdown("#### 📋 신고 상세")
            st.write(f"**자동 요약:** {selected_report.get('auto_summary', '')}")
            if selected_report.get("additional_note"):
                st.write(f"**사용자 추가 의견:** {selected_report.get('additional_note', '')}")
            st.text_area("최근 동작 순서", value=selected_report.get("action_history", ""), height=170,
                         disabled=True, key=f"admin_action_history_{selected_report_id}")
            with st.expander("기술 정보 보기"):
                st.json({
                    "error_type": selected_report.get("error_type", ""),
                    "error_message": selected_report.get("error_message", ""),
                    "last_input": selected_report.get("last_input", ""),
                    "response_time": selected_report.get("response_time"),
                    "screenshot": selected_report.get("screenshot", ""),
                    "traceback": selected_report.get("technical_traceback", ""),
                })


# =========================================================
# 상시 오류 신고 버튼 / AI 상태 표시 / 자동 오류 팝업
# =========================================================
render_floating_error_button()
render_ai_status()

if st.session_state.get("pending_auto_error_dialog"):
    st.session_state["pending_auto_error_dialog"] = False
    st.session_state.pop("error_report_auto_summary", None)
    st.session_state.pop("error_report_additional_note", None)
    show_error_report_dialog()
